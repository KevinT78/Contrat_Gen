"""Moteur de contrats : remplit les {{Placeholder}} d'un template.

Le template peut etre un .docx ou un .html ; le contrat PRODUIT est toujours un
.docx (un .html rempli est rendu en Word via htmldocx). Le circuit RH, la
remise au comptable et la signature ne manipulent donc qu'un seul format.

`paragraphes` et `remplacer` sont repris de wingstop_/contrat.py (le seul bloc
reellement reemployable, zero reseau) ; les regles de balisage calibrees sur la
prose Wingstop ne sont PAS reprises -- ici les templates arrivent deja balises.
"""
import io
import re
from datetime import date
from pathlib import Path

from docx import Document
from htmldocx import HtmlToDocx
# htmldocx passe le HTML entier a BeautifulSoup, qui croit y voir un nom de
# fichier des qu'il est court : avertissement sans objet, il pollue la console.
import warnings
warnings.filterwarnings("ignore", message=".*looks more like a filename.*")
from num2words import num2words

MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
           "août", "septembre", "octobre", "novembre", "décembre"]

# Placeholders que l'app calcule elle-meme (jamais saisis au formulaire).
CALCULES = ("Aujourdhui", "FaitLe", "NomPrenom", "NomNaissanceUsage", "DateSignature")


def paragraphes(doc):
    """Corps, tableaux, en-tetes et pieds de page -- chaque paragraphe une fois."""
    # dict et pas set d'id() : garder une reference vivante sur chaque paragraphe,
    # sinon lxml libere le proxy et reattribue son id() a un autre element, qui
    # est alors saute en silence.
    vus = {}

    def des(conteneur):
        for p in conteneur.paragraphs:
            if id(p._p) not in vus:
                vus[id(p._p)] = p
                yield p
        for t in conteneur.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from des(cell)

    yield from des(doc)
    for s in doc.sections:
        for partie in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            yield from des(partie)


def remplacer(p, motif, repl, limite=10):
    """Remplace dans un paragraphe en ne touchant que les runs couverts par le
    match : sinon Word perd le gras, l'italique et les sauts de ligne du reste."""
    n, depuis = 0, 0
    while n < limite:
        runs = p.runs
        texte = "".join(r.text for r in runs)
        m = motif.search(texte, depuis)
        if not m:
            return n
        nouveau, debut, fin = m.expand(repl), m.start(), m.end()
        if texte[debut:fin] == nouveau:
            return n                        # rien a changer, et surtout pas de boucle
        depuis = debut + len(nouveau)
        pos = 0
        for r in runs:
            d, f = pos, pos + len(r.text)
            pos = f
            if f <= debut or d >= fin:      # run entierement hors du match
                continue
            r.text = r.text[:max(0, debut - d)] + nouveau + (r.text[fin - d:] if f > fin else "")
            nouveau = ""                    # le reste du match est efface
        n += 1
    return n


def date_fr(valeur):
    """-> « 2 septembre 2026 » (1er pour le 1er du mois). Avale l'ISO nu
    (2026-09-02), l'ISO horodate (2026-09-02T07:00:00Z) et le jj/mm/aaaa.
    Rend la chaine telle quelle si rien ne matche -- jamais d'exception."""
    s = str(valeur or "")
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s) or re.match(r"(\d{2})/(\d{2})/(\d{4})\Z", s)
    if not m:
        return s
    a, mois, j = (m[1], m[2], m[3]) if "-" in s else (m[3], m[2], m[1])
    if not 1 <= int(mois) <= 12:
        return s
    return f"{'1er' if int(j) == 1 else int(j)} {MOIS_FR[int(mois) - 1]} {a}"


def sans_accent(s):
    return str(s or "").translate(str.maketrans(
        "àâäéèêëîïôöùûüçÀÂÄÉÈÊËÎÏÔÖÙÛÜÇ", "aaaeeeeiioouuucAAAEEEEIIOOUUUC"))


def prenom_nom(chaine):
    """« NOM Prenom » (formulaire) -> « Prenom NOM » (redaction des contrats).
    Ne reordonne QUE si des majuscules signalent le nom (« NKEMBA Awa » ->
    « Awa NKEMBA », « DE DOM Armand » -> « Armand DE DOM ») ; sinon rend tel
    quel -- on n'invente pas ou finit un nom compose non tirete."""
    mots = str(chaine or "").split()
    maj = [m for m in mots if len(m) > 1 and m == m.upper() and m.upper() != m.lower()]
    autres = [m for m in mots if m not in maj]
    if not maj or not autres:
        return str(chaine or "").strip()
    return " ".join(autres + maj)


def calculees(champs):
    """Le jeu fixe de valeurs derivees, en code produit (per ticket 05)."""
    import config
    prenom, nom = config.identite(champs)
    aujourdhui = date_fr(date.today().isoformat())
    return {
        "Aujourdhui": aujourdhui,
        "FaitLe": aujourdhui,
        "DateSignature": aujourdhui,
        # Champs separes -> « Prenom NOM ». Champ unique « NOM Prenom » (aucun
        # champ prenom dans le formulaire du client) -> reordonne pour la prose.
        # Le discriminant est l'absence de PRENOM, pas celle du nom : un client
        # sans champ « nom de naissance » a quand meme un nom, via le role nom.
        "NomPrenom": (f"{prenom} {nom.upper()}".strip() if prenom
                      else prenom_nom(nom)),
        # Meme cascade que identite() : un client sans champ « nom de
        # naissance » a quand meme un nom, via le role nom.
        "NomNaissanceUsage": (config.valeur(champs, "nom_naissance") or nom).upper(),
    }


# --- derivations declarees en config (formateurs + blocs conditionnels) -----
#
# Le pont generique client -> template : ni les {{Placeholders}} des modeles ni
# les champs du formulaire ne sont a adapter, c'est la config qui les relie.
#   {"placeholder": "X", "format": "lettres", "de": "Y"}   -> formateur integre
#   {"placeholder": "X", "si": ["champ","==","v"],          -> bloc conditionnel
#    "alors": "texte {{avec}} placeholders", "sinon": ""}

def _nombre(v):
    s = re.sub(r"[^\d,.-]", "", str(v))
    # « 1.867,02 » : le point separe les MILLIERS, la virgule les decimales.
    # Les laisser tous deux devenir « . » donnait « 1.867.02 » -> ValueError ->
    # 0 en silence, et un salaire en lettres « zero » dans le contrat signe.
    #
    # Sans virgule, « 1.480 » est ambigu. La FORME tranche : un point suivi
    # d'EXACTEMENT trois chiffres (repete) est un separateur de milliers
    # (« 1.480 » = 1480, « 1.480.500 » = 1480500) ; deux chiffres ou moins,
    # c'est un decimal (« 1480.08 », « 1.5 »). Sans cette regle « 1.480 »
    # devenait 1,48 -- la variante non-« zero » du meme degat.
    if "," in s or re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", s):
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        return float(s) if "." in s else int(s or 0)
    except ValueError:
        return 0


def montant_fr(v):
    """1867.05 -> « 1 867,05 » ; 2500 -> « 2 500 ». Format des contrats client :
    separateur de milliers = espace, ,00 supprime. (repris de wingstop_)"""
    entier, _, cents = f"{float(v):.2f}".partition(".")
    milliers = f"{int(entier):,}".replace(",", " ")
    return milliers if cents == "00" else f"{milliers},{cents}"


def lettres_fr(v):
    """Montant en toutes lettres pour la parenthese du contrat, JAMAIS
    « euro(s) » ni « centime(s) » -- c'est la prose du template qui ecrit
    l'unite (ex. « ({{SalaireLettres}}) euros bruts »). Sans centimes :
    cardinal seul (« deux mille »). Avec centimes : « <entier> virgule
    <centimes> » -- les centimes sont construits chiffre a chiffre plutot que
    lus depuis le float par num2words, qui confond 1867,50 et 1867,05 (0.5 et
    0.05 s'ecrivent tous deux « virgule cinq »)."""
    # Arrondir le MONTANT, pas les seuls centimes : round((n - entier) * 100)
    # sur 1867,999 donnait 100 centimes -> « virgule cent ». La retenue doit
    # remonter sur l'unite (-> « mille huit cent soixante-huit »).
    n = round(_nombre(v), 2)
    entier = int(n)
    centimes = round((n - entier) * 100)
    if centimes == 0:
        return num2words(entier, lang="fr")
    mot = num2words(centimes, lang="fr")
    if centimes < 10:
        mot = "zéro " + mot          # « 05 » se lit « zéro cinq », pas « cinq »
    return f"{num2words(entier, lang='fr')} virgule {mot}"


FORMATEURS = {
    "lettres": lettres_fr,
    "mensualise": lambda v: montant_fr(round(_nombre(v) * 52 / 12, 2)),  # h/sem -> h/mois
    "nombre": lambda v: str(_nombre(v)) if str(v).strip() else "",  # « 24H » -> « 24 »
    "majuscules": lambda v: str(v).upper(),
    "date_longue": lambda v: date_fr(str(v)),
}


def _mots(s):
    """Ensemble des mots normalises : « Equipier_Partiel » -> {equipier, partiel}."""
    return {w for w in re.split(r"[^a-z0-9]+", sans_accent(s).lower()) if w and w != "dk"}


def ligne_grille(poste, grille):
    """La ligne de grille qui remunere ce poste, ou None.

    Match par mots normalises, la plus specifique gagne : « Assistant Manager »
    ne prend pas le tarif « Manager », et « Equipier Polyvalent » prend celui
    d'« Equipier » quand le client n'a pas fait plus fin. Le revers de cette
    souplesse -- un poste sans ligne qui tombe sur celle d'un poste dont
    l'intitule est contenu dans le sien -- est refuse au demarrage par
    config._verifier_grille(), pas ici : a la generation, il est trop tard."""
    demande = _mots(poste)
    if not demande or not grille:
        return None
    compat = [e for e in grille.get("postes", []) if _mots(e["poste"]) <= demande]
    return max(compat, key=lambda e: len(_mots(e["poste"]))) if compat else None


def salaire(champs, grille):
    """(chiffres, lettres) de la remuneration mensuelle brute depuis la grille du
    client -- le formulaire n'en collecte pas.

    Poste au forfait  : {"poste": ..., "mensuel": N}      -> montant direct.
    Poste au SMIC     : {"poste": "Equipier", "champ_heures"?, "bareme":
                         {"<heures>": {"chiffres"}}}  -> ligne par duree hebdo
                         du contrat, "chiffres" PRIS TEL QUEL (aucun calcul ici
                         -> aucune divergence d'arrondi avec le contrat papier) ;
                         les lettres sont DERIVEES de ce meme "chiffres", pas
                         lues d'un second champ fige -- une seule source.
    Poste inconnu ou duree hors bareme -> ("", "") : le contrat refuse alors de
    sortir plutot que de porter un mauvais salaire. Match du poste par mots
    normalises, le plus specifique gagne (« Assistant Manager » ne prend pas le
    tarif « Manager »)."""
    import config
    e = ligne_grille(config.valeur(champs, "poste"), grille)
    if e is None:
        return "", ""
    if "mensuel" in e:
        v = float(e["mensuel"])
        return montant_fr(v), lettres_fr(v)
    ligne = e.get("bareme", {}).get(cle_bareme(champs.get(grille.get("champ_heures", ""), "")))
    return (ligne["chiffres"], lettres_fr(ligne["chiffres"])) if ligne else ("", "")


def cle_bareme(duree):
    """« 24H » -> "24" : la cle de bareme d'une duree hebdo, "" sans chiffre.
    Partagee avec config._verifier_bareme : le garde-fou lit la duree
    exactement comme la generation du contrat."""
    mo = re.match(r"\d+", str(duree))
    return mo.group() if mo else ""


OPERATEURS = ("==", "!=", "in", "present", "absent")


def _teste(triple, ctx):
    if not triple:
        return True
    champ, op, val = triple
    g = str(ctx.get(champ, ""))
    return {"==": g == str(val), "!=": g != str(val), "in": g in (val or []),
            "present": bool(g), "absent": not g}[op]


def _expanse(texte, d):
    return re.sub(r"\{\{(\w+)\}\}",
                  lambda m: str(d.get(m.group(1), m.group(0))), texte or "")


def _appliquer_derives(vals, champs):
    import config
    regles = config.derives()
    if not regles:
        return vals
    ctx = {c["id"]: champs.get(c["id"], "") for c in config.champs()}
    ctx.update(vals)
    for d in regles:
        ph = d["placeholder"]
        if "format" in d:
            vals[ph] = FORMATEURS[d["format"]](ctx.get(d["de"], ""))
        else:
            texte = d["alors"] if _teste(d.get("si"), ctx) else d.get("sinon", "")
            vals[ph] = _expanse(texte, ctx)
        ctx[ph] = vals[ph]
    return vals


def valeurs(champs, mentions, extra=None):
    """champs (ids du schema) + mentions societe/etablissement -> placeholders.

    `extra` : placeholders calcules par l'appelant qui ne sont ni des champs ni
    des mentions (ex. {{PiecesFournies}} / {{PiecesManquantes}} de la fiche
    salarie, qui dependent de l'etat du disque)."""
    import config

    vals = dict(mentions)
    for c in config.champs():
        ph = c.get("placeholder")
        if not ph:
            continue
        v = champs.get(c["id"], "")
        vals[ph] = date_fr(v) if c["type"] == "date" else str(v)
    vals.update(calculees(champs))
    # Salaire depuis la grille du client (le formulaire n'en collecte pas). Une
    # saisie RH dans `extra` reste prioritaire -- correction manuelle possible.
    grille = config.grille()
    if grille:
        ch, le = salaire(champs, grille)
        if ch:
            vals["SalaireChiffres"], vals["SalaireLettres"] = ch, le
        else:
            vals.setdefault("SalaireChiffres", "")
            vals.setdefault("SalaireLettres", "")
    vals.update(extra or {})
    _appliquer_derives(vals, champs)
    return vals


def elaguer(doc, vals):
    """Fiche generique : retire chaque ligne de tableau dont AUCUN {{jeton}}
    n'a de valeur ici (absent de la config du client, ou reponse vide -- ex.
    titre de sejour d'un salarie francais), puis tout tableau devenu vide avec
    son titre de section juste au-dessus. Plusieurs jetons dans une ligne =
    synonymes d'une config a l'autre ({{NumSS}}{{NumeroSecu}}) : le premier
    alimente s'affiche, les autres sont vides par _generer_docx. Retourne les
    libelles retires (1re cellule). Tableaux du corps seulement : c'est la
    forme du modele."""
    retires = []
    for t in list(doc.tables):
        for row in list(t.rows):
            jetons = re.findall(r"\{\{(\w+)\}\}", "".join(c.text for c in row.cells))
            if jetons and not any(str(vals.get(j, "")).strip() for j in jetons):
                retires.append(row.cells[0].text)
                row._tr.getparent().remove(row._tr)
        if not t.rows:
            titre = t._tbl.getprevious()
            if titre is not None and titre.tag.endswith("}p"):
                titre.getparent().remove(titre)
            t._tbl.getparent().remove(t._tbl)
    return retires


def generer(template, vals, dest=None, elaguer_lignes=False, retires=None):
    """Remplit les {{Placeholder}}. Refuse d'ecrire s'il en reste un : un contrat
    troue -- ou pire, portant le nom de l'ancien salarie -- ne doit pas sortir.
    Le produit est TOUJOURS un .docx : template .docx via python-docx (runs
    preserves), template .html (ou .txt/.md) en substitution texte plate puis
    rendu Word via htmldocx.

    `dest` None -> retourne les octets du document, a charge de l'appelant de
    les confier a store.poser_octets (le moteur ne connait pas le stockage).
    `dest` fourni -> ecrit dans ce Path et le retourne (doctor, tests).
    `elaguer_lignes` (modele .docx seulement) : voir elaguer().
    `retires` liste fournie -> les libelles elagues y sont ajoutes. Le retour
    est deja pris (octets ou Path) et doctor, seul interesse, rouvrait le .docx
    pour rejouer elaguer() sur ce que la generation venait de calculer."""
    # python-docx leve PackageNotFoundError (pas un OSError) sur un fichier
    # absent : doctor et app ne l'attrapaient pas et plantaient au lieu de
    # marquer le cas « refuse ».
    if not Path(template).exists():
        raise FileNotFoundError(f"modèle absent : {template}")
    if Path(template).suffix.lower() == ".docx":
        return _generer_docx(template, vals, dest, elaguer_lignes, retires)
    return _texte_vers_docx(remplir(template, vals), dest)


def remplir(template, vals):
    """Template texte (.html/.txt/.md) -> son contenu, {{jetons}} substitues.
    Leve si une valeur critique est vide ou si un {{jeton}} reste orphelin.
    Point d'entree du rendu (generer) et des tests qui asserten la prose."""
    txt = Path(template).read_text(encoding="utf-8")
    _verifier_critiques(Path(template).name, txt, vals)
    for cle, val in vals.items():
        txt = txt.replace("{{" + cle + "}}", str(val))
    restants = _restants(txt)
    if restants:
        raise ValueError(f"placeholders non remplis dans {Path(template).name} : "
                         + ", ".join(restants))
    return txt


def _texte_vers_docx(html, dest):
    """HTML rempli -> .docx via htmldocx. `dest` None -> octets ; sinon ecrit."""
    doc = Document()
    HtmlToDocx().add_html_to_document(html, doc)
    if dest is None:
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest


def _restants(texte):
    return sorted(set(re.findall(r"\{\{[^}\n]{0,40}\}\}", texte)))


def _verifier_critiques(nom, texte_brut, vals):
    """Un placeholder critique (champ requis, ou liste instance['critiques'])
    remplace par du VIDE ne laisse aucun {{}} a detecter -- le contrat sort avec
    une ligne de paie blanche, sans signal (bug wingstop_ du 2026-08-25). On
    refuse avant substitution, sur les placeholders reellement dans le template."""
    import config
    utilises = set(re.findall(r"\{\{(\w+)\}\}", texte_brut))
    vides = sorted(c for c in config.critiques()
                   if c in utilises and not str(vals.get(c, "")).strip())
    if vides:
        raise ValueError(f"valeurs vides dans {nom} : " + ", ".join(vides))


def _generer_docx(template, vals, dest, elaguer_lignes=False, retires=None):
    doc = Document(str(template))
    # critiques AVANT l'elagage : une valeur requise vide doit refuser la
    # fiche, pas faire disparaitre sa ligne en silence.
    _verifier_critiques(Path(template).name,
                        "\n".join(p.text for p in paragraphes(doc)), vals)
    if elaguer_lignes:
        coupes = elaguer(doc, vals)
        if retires is not None:
            retires.extend(coupes)
        # Synonymes non alimentes des lignes GARDEES -> vides, pas « troues ».
        # Restreint aux cellules de tableau : c'est la seule zone qu'elaguer
        # sait couper. Blanchir tout le document faisait sortir en silence un
        # jeton de prose ou d'en-tete que la config n'alimente pas, au lieu de
        # le refuser comme dans n'importe quel autre modele.
        cellules = "\n".join(c.text for t in doc.tables
                             for row in t.rows for c in row.cells)
        vals = {**{j: "" for j in re.findall(r"\{\{(\w+)\}\}", cellules)}, **vals}
    for p in paragraphes(doc):
        for cle, val in vals.items():
            remplacer(p, re.compile(r"\{\{" + re.escape(cle) + r"\}\}"),
                      str(val).replace("\\", r"\\"))
    restants = _restants("\n".join(p.text for p in paragraphes(doc)))
    if restants:
        raise ValueError(f"placeholders non remplis dans {Path(template).name} : "
                         + ", ".join(restants))
    if dest is None:
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest
