"""Configuration du client : tout ce qui varie d'un client à l'autre vit en
fichiers sous config/, jamais en code.

Un client = une copie de ce dépôt. `python installer.py "<Client>"` remet
config/ à blanc et écrit les secrets dans un fichier d'environnement ; on
remplit ensuite instance.json, les sociétés en base (configurer.py) et on
dépose les .docx. Les secrets ne vivent jamais dans config/ (cf. SECRETS).
"""
import json
import os
import re
import sys
import time
from pathlib import Path

from pymongo.errors import DuplicateKeyError, PyMongoError

from apparence import resoudre

RACINE = Path(__file__).parent


def _est_env(a):
    return a == "--env" or a.startswith("--env=")


def _charger_env(argv):
    """`--env <fichier>` (ou `--env=<fichier>`) sur la ligne de n'importe quel
    script (app, recap, configurer, doctor...) : charge le .env de l'instance
    ecrit par installer.py. Ici parce que config est importe avant tout le reste
    et que CONFIG_DIR doit etre connu avant CLIENT ; l'option est retiree d'argv,
    chaque script lit ses propres arguments comme avant.

    Une variable deja posee et non vide l'emporte sur le fichier (systemd,
    PowerShell, surcharge ponctuelle). Une ligne = CLE=valeur coupee au premier
    « = », espaces autour retires, une paire de guillemets englobante retiree,
    rien d'autre : la meme lecture sous Linux et Windows. Une ligne illisible
    arrete tout, sans jamais afficher son contenu (elle peut porter un secret)."""
    i = next((i for i, a in enumerate(argv) if _est_env(a)), None)
    if i is None:
        return
    if argv[i] == "--env":
        if i + 1 >= len(argv):
            raise SystemExit("--env attend le chemin du fichier .env de l'instance")
        chemin = Path(argv[i + 1])
        del argv[i:i + 2]
    else:
        chemin = Path(argv.pop(i).split("=", 1)[1])
    if any(_est_env(a) for a in argv):
        raise SystemExit("--env donné deux fois : un seul fichier par instance")
    try:
        # utf-8-sig : un .env retouche dans le Bloc-notes peut commencer par un BOM
        lignes = chemin.read_text(encoding="utf-8-sig").splitlines()
    except OSError as e:
        raise SystemExit(f"--env : {chemin} illisible ({e.strerror})")
    except UnicodeDecodeError:
        raise SystemExit(f"--env : {chemin} n'est pas en UTF-8, réenregistrez-le en UTF-8")
    for n, ligne in enumerate(lignes, 1):
        if not ligne.strip() or ligne.lstrip().startswith("#"):
            continue
        cle, egal, valeur = (x.strip() for x in ligne.partition("="))
        if not egal or not cle:
            raise SystemExit(f"--env : {chemin}, ligne {n} : attendu CLE=valeur")
        if len(valeur) >= 2 and valeur[0] == valeur[-1] and valeur[0] in "\"'":
            valeur = valeur[1:-1]
        if not os.environ.get(cle):
            os.environ[cle] = valeur
        elif os.environ[cle] != valeur:
            # La base d'un autre client restee dans le shell recevrait la
            # config de celui-ci. Le nom seul : la valeur peut etre un secret.
            print(f"! --env : {cle} déjà posé dans l'environnement avec une autre "
                  f"valeur que {chemin} — c'est celle de l'environnement qui sert",
                  file=sys.stderr)


_charger_env(sys.argv)
CLIENT = Path(os.environ.get("CONFIG_DIR", RACINE / "config"))

# Lus dans l'environnement, jamais dans config/ ni en base : config/ se copie,
# se versionne et se montre au client ; un secret n'y a pas sa place.
SECRETS = {"SECRET_KEY": "clé de session et de signature des liens",
           "MONGO_URI": "adresse de la base", "MONGO_DB": "nom de la base de ce client"}
MDP_PAR_DEFAUT = "demo"          # l'ancien mot de passe de démo — refusé au démarrage

# Version du FORMAT de config/, pas du produit. Le code est un paquet versionne,
# config/ vit chez le client : une config ecrite pour un format qu'on ne
# comprend plus doit refuser de demarrer en disant quoi faire, pas se faire
# reecrire dans le dos.
CONFIG_VERSION = 1

# Cache simple plutot que lru_cache : `recharger()` a besoin de reposer
# l'ancienne config si la nouvelle est invalide, ce qu'un cache_clear() ne
# permet pas.
_CACHE = {}


def _lire(nom):
    if nom not in _CACHE:
        _CACHE[nom] = json.loads((CLIENT / nom).read_text(encoding="utf-8"))
    return _CACHE[nom]


def instance():
    return _lire("instance.json")


def formulaire():
    return _lire("formulaire.json")


_SOCIETES_TTL = 60
_societes_lu = 0.0


def societes():
    """Societes et etablissements : la collection `societes`, triee par `ordre`,
    sans `_id` ni `ordre` : [{nom, siren, comptable_email, mentions,
    etablissements: [{nom, siret, groupe?, manager_email, contrat?, mentions}]}].

    Cache de _SOCIETES_TTL secondes, vide a chaque ecriture de l'app. Le delai
    borne deux retards : une lecture lancee avant une ecriture qui repose
    l'ancienne liste apres l'invalidation, et un import fait par un autre
    process (configurer.py) pendant que l'app tourne."""
    global _societes_lu
    if "societes" not in _CACHE or time.monotonic() - _societes_lu > _SOCIETES_TTL:
        import store      # store importe config : import local, cycle sinon
        _CACHE["societes"] = [
            {k: v for k, v in s.items() if k != "ordre"}
            for s in store.db().societes.find({}, {"_id": 0}).sort("ordre", 1)]
        _societes_lu = time.monotonic()
    return _CACHE["societes"]


def importer_societes(liste):
    """Charge une liste de societes (la forme que rend societes()). Chaque societe
    remplace celle de meme nom, les autres restent : rejouable, et rien n'est
    efface de ce qui n'est pas dans le fichier. Une societe deja en base garde
    sa place, une nouvelle se range a la suite. `_id` (export Mongo) est ignore."""
    import store
    col = store.db().societes
    places = {d["nom"]: d.get("ordre", 0) for d in col.find({}, {"nom": 1, "ordre": 1})}
    suivante = max(places.values(), default=-1) + 1
    for s in liste:
        if s["nom"] not in places:
            places[s["nom"]], suivante = suivante, suivante + 1
        doc = {k: v for k, v in s.items() if k != "_id"}
        col.replace_one({"nom": s["nom"]}, {**doc, "ordre": places[s["nom"]]}, upsert=True)
    _CACHE.pop("societes", None)


def sauver_etablissement(origine, societe_nom, societe, etablissement):
    """Cree ou modifie un etablissement, en base.

    `origine` vaut None a la creation, sinon (nom de societe, nom d'etablissement)
    avant modification. `societe` vaut None quand on rattache a une societe
    deja la sans toucher a sa fiche ; sinon c'est la fiche a ecrire
    (nom, siren, comptable_email, mentions). ValueError : nom deja pris,
    societe ou etablissement introuvable.
    """
    import store
    fiche = None
    if societe is not None:
        nom = (societe.get("nom") or "").strip()
        if not nom:
            raise ValueError("Le nom de la société est obligatoire.")
        # Un renommage vers le nom d'une autre societe creerait deux fiches
        # homonymes : la cle « Societe / Etablissement » deviendrait ambigue.
        if origine is not None and nom != origine[0] and any(
                s["nom"] == nom for s in societes()):
            raise ValueError("Cette société existe déjà : choisissez un autre nom.")
        fiche = {"nom": nom, "siren": societe.get("siren") or "",
                 "comptable_email": societe.get("comptable_email") or "",
                 "mentions": societe.get("mentions") or {}}
    col = store.db().societes
    try:
        if origine is None:
            _ajouter_etablissement(col, societe_nom, fiche, etablissement)
        else:
            _modifier_etablissement(col, origine, fiche, etablissement)
    except DuplicateKeyError:       # course avec un autre ecrivain : index unique `nom`
        raise ValueError("Cette société existe déjà : choisissez un autre nom.") from None
    finally:
        _CACHE.pop("societes", None)


def _ajouter_etablissement(col, societe_nom, fiche, etablissement):
    cible = col.find_one({"nom": societe_nom})
    if cible is None:
        if not fiche:
            raise ValueError("Société inconnue.")
        dernier = col.find_one(sort=[("ordre", -1)])
        col.insert_one({**fiche, "ordre": (dernier or {}).get("ordre", -1) + 1,
                        "etablissements": [etablissement]})
    elif fiche:
        raise ValueError("Cette société existe déjà : choisissez-la dans la liste.")
    elif not col.update_one(
            {"_id": cible["_id"], "etablissements.nom": {"$ne": etablissement["nom"]}},
            {"$push": {"etablissements": etablissement}}).matched_count:
        raise ValueError("Cet établissement existe déjà.")


def _modifier_etablissement(col, origine, fiche, etablissement):
    ancien_soc, ancien_etab = origine
    cible = col.find_one({"nom": ancien_soc})
    if cible is None:
        raise ValueError("Société introuvable.")
    etabs = cible.get("etablissements") or []
    idx = next((i for i, e in enumerate(etabs) if e["nom"] == ancien_etab), None)
    if idx is None:
        raise ValueError("Établissement introuvable.")
    if any(e["nom"] == etablissement["nom"] and i != idx for i, e in enumerate(etabs)):
        raise ValueError("Cet établissement existe déjà.")
    # Filtre sur le nom a cet indice : si la liste a bouge entre la lecture et
    # l'ecriture, on ne remplace pas un autre etablissement.
    if not col.update_one(
            {"_id": cible["_id"], f"etablissements.{idx}.nom": ancien_etab},
            {"$set": {f"etablissements.{idx}": etablissement, **(fiche or {})}}
    ).matched_count:
        raise ValueError("Établissement introuvable.")


def comptes():
    """Les comptes de connexion : la collection `comptes`, rendue sous la forme
    {identifiant: {"mdp_hash": ...}}. Jamais dans config/ : un hash de mot de
    passe n'est pas du parametrage."""
    import store          # store importe config : import local, cycle sinon
    return {d["_id"]: {"mdp_hash": d.get("mdp_hash")} for d in store.db().comptes.find()}


# --- la DA du client ------------------------------------------------------
# Un client = un dossier d'instance ; sa direction artistique y vit comme le
# reste (`theme.json` + un dossier `marque/` pour le logo et le favicon). Les
# regles, elles, vivent dans apparence.py : c'est du calcul de couleur et de
# la validation, pas de la lecture de config.
#
# MARQUE est servi par Flask lui-meme (static_folder), pas par une route
# maison : mimetypes, cache et protection de chemin sont deja faits.
MARQUE = CLIENT / "marque"


def theme():
    """`theme.json` de l'instance. Absent = la DA d'origine, une instance
    installee avant ce fichier ne bouge pas.

    Garde sur OSError SEULEMENT, pas sur le JSON casse comme conservation() : un
    fichier illisible avale la DA en silence, et le client ouvre son app le
    premier jour, la voit verte, sans que personne sache pourquoi. C'est
    _verifier_theme() qui le dit au demarrage."""
    try:
        return _lire("theme.json")
    except OSError:
        return {}


def apparence():
    """La DA du client, prete a poser dans base.html. Lecture GARDEE : un
    theme.json invalide a deja empeche l'instance de demarrer."""
    return resoudre(theme(), MARQUE).tokens


def conservation():
    """Le bloc `conservation` d'instance.json : au bout de combien de jours dans
    un etat terminal les pieces d'un dossier sont effacees (le journal, lui,
    reste). Lecture GARDEE : config est importe par des outils qui tournent
    sans instance complete (doctor, tests). Bloc absent = aucune purge."""
    try:
        c = instance().get("conservation")
    except (OSError, ValueError, KeyError):
        return {}
    return c if isinstance(c, dict) else {}


def compression():
    """`compression` d'instance.json : "couleur" (défaut) ou "gris". Pièces
    déposées (formulaire, RH) et importées de SharePoint passent par
    compresseur.py. Lecture GARDEE comme conservation()."""
    try:
        c = instance().get("compression")
    except (OSError, ValueError, KeyError):
        return "couleur"
    return c if c in ("couleur", "gris") else "couleur"


def sharepoint():
    """Le bloc `sharepoint` d'instance.json : le site d'ou le bouton
    « Récupérer les dossiers » importe les salaries deja en poste. Lecture
    GARDEE comme conservation(). Bloc absent = pas de bouton."""
    try:
        c = instance().get("sharepoint")
    except (OSError, ValueError, KeyError):
        return {}
    return c if isinstance(c, dict) else {}


def recharger():
    """Relit config/ a chaud, sans redemarrer. Renvoie {} si la nouvelle config
    est valide (elle est alors active), sinon les manques -- et l'ancienne
    config reste en place : une instance qui servait continue de servir."""
    ancien = dict(_CACHE)
    _CACHE.clear()
    if manques := verifier():
        _CACHE.clear()
        _CACHE.update(ancien)
        return manques
    return {}


def etablissements():
    """[(cle, libelle)] pour le select du formulaire. Cle = 'Societe / Etab'."""
    return [(f"{s['nom']} / {e['nom']}", f"{s['nom']} — {e['nom']}")
            for s in societes() for e in s["etablissements"]]


def etablissement(cle):
    """Retrouve (societe, etablissement) depuis la cle du formulaire."""
    for s in societes():
        for e in s["etablissements"]:
            if f"{s['nom']} / {e['nom']}" == cle:
                return s, e
    raise KeyError(f"etablissement inconnu : {cle}")


def mentions(cle):
    """Placeholders legaux de la societe + de l'etablissement, pour le contrat."""
    s, e = etablissement(cle)
    return {**s.get("mentions", {}), **e.get("mentions", {}),
            "Siret": e.get("siret", ""), "Siren": s.get("siren", ""),
            "Etablissement": e["nom"], "Societe": s["nom"]}


def comptable(cle):
    s, _ = etablissement(cle)
    return s.get("comptable_email") or instance()["mails"]["comptable_defaut"]


def groupe(cle):
    """Tete de l'arborescence des dossiers salaries : la societe ->
    etablissement.groupe, a defaut le nom de la societe."""
    s, e = etablissement(cle)
    return e.get("groupe") or s["nom"]


def manager(cle):
    """Adresse fixe du manager de l'etablissement (en base), ou "" --
    y compris quand l'etablissement a disparu de la config (meme regle que
    mode_contrat) : valider ou rejeter un vieux dossier ne doit pas 500."""
    try:
        _, e = etablissement(cle)
    except KeyError:
        return ""
    return e.get("manager_email", "")


def url_publique():
    """Base de TOUS les liens des mails (app._lien, recap.py) : instance.json -> url,
    jamais l'hote de la requete."""
    return (instance().get("url") or "").rstrip("/")


def mode_contrat(cle):
    """Couloir du schema, en donnee : 'genere' (Dark Kitchen) ou 'depose'
    (Restaurant, contrat fait a la main dans le logiciel RH). Absent = 'genere' -- y
    compris quand l'etablissement lui-meme a disparu de la config (dossier
    vieux ou purge) : un ecran dossier ne doit pas 500 pour ca."""
    try:
        _, e = etablissement(cle)
    except KeyError:
        return "genere"
    return e.get("contrat", "genere")


def _normaliser(x):
    """Comparaison de valeurs de config insensible a la casse et aux accents :
    « Equipier polyvalent » exporte par MS Forms doit matcher « Equipier
    Polyvalent », « Oui » matcher « OUI ». Import local de sans_accent : c'est
    contrat qui importe config, pas l'inverse."""
    from contrat import sans_accent
    return sans_accent(x).casefold().strip()


def regle_pour(champs):
    """La regle de template qui matche ces champs, ou None si aucune ne colle.

    Meme matching que modele_pour, expose a part pour distinguer « aucune
    regle ne vise ce cas » (regle_pour renvoie None) de « regle d'exclusion
    volontaire » (regle trouvee, mais `modele: null` -- un cas identifie sans
    contrat a produire, ex. « Manager » a temps partiel qui n'existe pas chez
    ce client) sans dupliquer la comparaison. Voir modele_pour pour les
    formes acceptees.
    """
    n = _normaliser
    for regle in regles():
        if all(n(champs.get(k, "")) == n(v)
               for k, v in regle.get("quand", {}).items()):
            return regle
    return None


def modele_pour(champs):
    """Poste (+ conditions eventuelles) -> nom du modele de contrat, ou None
    (aucune regle ne vise ce cas, ou regle d'exclusion volontaire -- voir
    regle_pour pour distinguer les deux ; ici le contrat doit refuser dans
    les deux cas).

    Deux formes acceptees pour instance()['templates'] :
      - dict {poste: fichier}                              (client simple)
      - liste [{"quand": {champ: valeur, ...}, "modele": fichier}, ...]
        premiere regle dont TOUS les champs 'quand' collent ; une regle sans
        'quand' est un fourre-tout. Sert au branchement temps partiel, ou a un
        modele par etablissement quand les mentions sont figees dans la prose.
        `modele: null` marque une exclusion volontaire (cas identifie, pas de
        contrat pour ce croisement) plutot que d'omettre la regle -- doctor
        l'affiche a part, distinct d'un cas oublie.

    Comparaison insensible a la casse et aux accents : « Equipier polyvalent »
    exporte par MS Forms matche « Equipier Polyvalent », « Oui » matche « OUI ».
    """
    regle = regle_pour(champs)
    return regle.get("modele") if regle else None


def derives():
    """Placeholders calcules par regle en config (formateurs, blocs conditionnels)."""
    return instance().get("derives", [])


def saisie_rh():
    """Placeholders qu'aucune question du formulaire ne fournit : la RH les
    saisit a la generation du contrat. Noms seuls ; voir saisie_rh_champs."""
    return [c["placeholder"] for c in saisie_rh_champs()]


def saisie_rh_champs():
    """Chaque entree de saisie_rh est soit "Placeholder", soit
    {"placeholder", "libelle", "aide"} -> toujours rendue sous la forme dict,
    libelle par defaut = nom du placeholder."""
    out = []
    for c in instance().get("saisie_rh", []):
        if isinstance(c, str):
            c = {"placeholder": c}
        out.append({"placeholder": c["placeholder"],
                    "libelle": c.get("libelle") or c["placeholder"],
                    "aide": c.get("aide", "")})
    return out


def critiques():
    """Placeholders qui, vides, produisent un contrat sans objet (ligne de paie
    blanche, « demeurant au . »). Derives des champs requis + liste explicite
    instance()['critiques'] (pour ce qui ne vient pas d'un champ, ex. le salaire
    tire de la grille)."""
    ph = {c["placeholder"] for c in champs()
          if c.get("requis") and c.get("placeholder")}
    return ph | set(instance().get("critiques", []))


def grille():
    """Grille de remuneration du client (salaire par poste). {} si absente : un
    client peut faire saisir le salaire a la RH plutot que le tirer d'une grille."""
    if "grille.json" not in _CACHE:
        f = CLIENT / "grille.json"
        _CACHE["grille.json"] = json.loads(f.read_text(encoding="utf-8")) if f.exists() else {}
    return _CACHE["grille.json"]


def champs():
    return formulaire()["champs"]


# Le vocabulaire ferme des roles que le MOTEUR sait lire, et l'id historique de
# chacun. Tout le reste du formulaire est du champ libre : le moteur l'ignore et
# il ne va qu'aux templates. C'est ce qui rend le produit installable chez une
# PME dont les champs s'appellent autrement -- avant, « etablissement » et
# « poste » etaient lus en dur a 13 endroits et un renommage levait un KeyError
# en pleine action RH.
ROLES = {"etablissement": "etablissement", "poste": "poste",
         "nom": "nom_naissance", "prenom": "prenom", "nom_usage": "nom_usage",
         "nom_naissance": "nom_naissance", "email": "email_demandeur",
         "date_debut": "date_debut"}

# Sans eux le parcours ne peut pas tourner : l'etablissement route les mentions
# et le comptable, le poste route le modele et la grille, le nom et l'email
# adressent les mails. Les autres sont facultatifs (un formulaire a champ unique
# « NOM Prenom » n'a ni prenom ni nom d'usage).
ROLES_REQUIS = ("etablissement", "poste", "nom", "email")


def role(nom):
    """id du champ qui joue un role attendu par le code. Declare dans
    formulaire()['roles'] ; a defaut l'id historique de ROLES -- une config sans
    'roles' garde le comportement d'avant.

    Rend TOUJOURS une chaine, meme sur une config malformee : `roles` qui n'est
    pas un objet, ou un id qui est une liste. Sinon chaque appelant explose a sa
    facon (`x not in ids` -> TypeError) -- y compris verifier(), dont c'est le
    travail de refuser cette config proprement."""
    declares = formulaire().get("roles") or {}
    declare = declares.get(nom) if isinstance(declares, dict) else None
    return declare if isinstance(declare, str) and declare else ROLES.get(nom) or nom


def valeur(champs, nom, defaut=""):
    """La valeur qu'un dossier porte pour un role. LE seul acces du moteur aux
    champs d'un salarie : aucun id de champ n'est ecrit en dur ailleurs."""
    return champs.get(role(nom), defaut)


def identite(champs):
    """(prenom, nom) d'un dossier, quel que soit le decoupage du formulaire du
    client : champs separes, ou champ unique « NOM Prenom » (-> ("", "NOM
    Prenom")). Les trois modules qui affichaient un nom le construisaient chacun
    de leur cote, et aucun ne retombait sur le role « nom » : un client sans
    champ « nom de naissance » sortait des contrats au prenom seul."""
    nom = (valeur(champs, "nom_usage") or valeur(champs, "nom_naissance")
           or valeur(champs, "nom"))
    return valeur(champs, "prenom"), nom


def champ(cid):
    for c in champs():
        if c["id"] == cid:
            return c
    # cle hors schema (ex. une saisie RH fusionnee dans champs) : libelle = la
    # cle elle-meme, plutot qu'un KeyError qui casse l'ecran dossier.
    return {"id": cid, "libelle": cid, "type": "texte"}


def pieces():
    return [c for c in champs() if c["type"] == "piece_jointe"]


def secret():
    """Cle de signature des liens (store.signer). Leve si SECRET_KEY manque :
    recap.py ne passe pas par verifier(), et une cle vide y signerait des liens
    morts sans que personne le voie."""
    if not (cle := os.environ.get("SECRET_KEY", "")).strip():
        raise RuntimeError("SECRET_KEY absent de l'environnement")
    return cle.encode()


# Modele generique verse au depot (hors config*) : sert a TOUTE instance qui
# n'en declare pas un a elle -- la fiche salarie est desormais toujours
# produite, plus besoin de la deposer pour l'avoir.
MODELE_FICHE_GENERIQUE = RACINE / "modeles" / "fiche_salarie.docx"


def fiche_salarie():
    """(chemin du modele de fiche salarie, est-ce le generique ?).

    Celui du CLIENT (config/contrats/) si instance()['fiche_salarie'] le
    declare, sinon le modele generique. Le booleen voyage AVEC le chemin
    parce qu'il en decoule : seul le generique s'elague (contrat.elaguer),
    un modele client est cense etre taille pour sa config. Chaque appelant
    le recalculait de son cote a partir de instance()."""
    if cle := instance().get("fiche_salarie"):
        return CLIENT / "contrats" / cle, False
    return MODELE_FICHE_GENERIQUE, True


def _templates_actifs():
    """Les modeles de CONTRAT (par poste ou par regle) dont un placeholder
    manquant doit bloquer le demarrage -- noms relatifs a config/contrats/.
    La fiche salarie est verifiee a part (_verifier_placeholders) : son
    modele generique ne vit pas forcement sous config/contrats/."""
    return {r["modele"] for r in regles() if r.get("modele")}


def regles():
    """Les regles `templates`, toujours en liste : la forme dict {poste: fichier}
    n'est qu'une ecriture abregee de [{"quand": {poste: ...}, "modele": ...}],
    convertie ici pour que matching et verifications ne la connaissent pas.
    Les regles mal formees sont sautees : _verifier_regles les refuse, et les
    autres lecteurs (doctor tourne meme sur une config refusee) ne doivent pas
    planter dessus."""
    t = instance()["templates"]
    if isinstance(t, dict):
        return [{"quand": {role("poste"): p}, "modele": m} for p, m in t.items()]
    return [r for r in t if isinstance(r, dict)
            and isinstance(r.get("quand", {}), dict)] if isinstance(t, list) else []


# Volontairement laxiste : on veut VOIR « {{ Nom }} » et « {{nom}} » pour les
# refuser au demarrage. Une regex stricte (\w+) les rend invisibles au garde-fou
# et les laisse exploser a la generation du contrat d'un vrai salarie -- c'est
# le mode d'echec a eviter quand c'est le CLIENT qui balise ses .docx.
JETON = re.compile(r"\{\{([^}\n]{0,60})\}\}")


def placeholders_connus():
    """{placeholder: d'ou vient sa valeur} pour CETTE config. Deux usages :
    le garde-fou de demarrage, et la fiche remise au client qui balise ses
    propres .docx (`python placeholders.py`)."""
    import contrat

    src = {}
    for c in champs():
        if c.get("placeholder"):
            src[c["placeholder"]] = f"Question « {c['libelle']} » du formulaire"
    for p in contrat.CALCULES:
        src[p] = "Calculé automatiquement"
    src["PiecesFournies"] = "Liste des pièces jointes fournies"
    src["PiecesManquantes"] = "Liste des pièces jointes manquantes"
    for d in derives():
        # .get : une regle sans placeholder est une config invalide, que
        # verifier() rapporte -- elle ne doit pas faire planter le rapport.
        if ph := d.get("placeholder"):
            src[ph] = "Règle dérivée (instance.json → derives)"
    for p in saisie_rh():
        src[p] = "Saisi par la RH au moment de générer le contrat"
    if grille():
        src["SalaireChiffres"] = "Grille de salaires (config/grille.json)"
        src["SalaireLettres"] = "Grille de salaires, en toutes lettres"
    for s in societes():
        for p in s.get("mentions", {}):
            src[p] = "Mention légale de la société (en base)"
        for e in s["etablissements"]:
            for p in e.get("mentions", {}):
                src[p] = "Mention légale de l'établissement (en base)"
    for p in ("Societe", "Siren", "Etablissement", "Siret"):
        src[p] = "Identité de la société / de l'établissement choisi"
    return src


def _jetons(chemin):
    """Les {{jetons}} bruts d'un template, tels qu'ecrits -- espaces compris."""
    if chemin.suffix.lower() == ".docx":
        from docx import Document
        import contrat
        texte = "\n".join(p.text for p in contrat.paragraphes(Document(chemin)))
    else:
        texte = chemin.read_text(encoding="utf-8")
    return set(JETON.findall(texte))


def _verifier_version():
    v = instance().get("config_version", CONFIG_VERSION)
    if v != CONFIG_VERSION:
        return {"config_version": [
            f"config/ est au format {v}, ce code attend le format "
            f"{CONFIG_VERSION} — mettez à jour config/instance.json "
            f"(voir CHANGELOG) ou réinstallez la version correspondante"]}
    return {}


def _verifier_mongo():
    """Les variables d'environnement, sans toucher au reseau : une base
    injoignable se voit au `ping` de app.demarrer(). Jamais la valeur de
    MONGO_URI dans le message, elle porte le mot de passe de la base."""
    absents = [v for v in SECRETS if not os.environ.get(v)]
    if absents:
        return {"mongo / secrets": [f"{v} absent de l'environnement ({SECRETS[v]}) — "
                                    "passez le .env de l'instance : --env <dossier>/.env"
                                    for v in absents]}
    raisons = []
    if len(os.environ["SECRET_KEY"].strip()) < 32:
        raisons.append("SECRET_KEY trop courte (moins de 32 caractères) : gardez celle "
                       "générée par installer.py")
    nom = os.environ["MONGO_DB"]
    # Les regles de nommage de MongoDB : un nom invalide fait lever pymongo au
    # premier acces, en pleine verification.
    if re.search(r'[/\\. "$*<>:|?]', nom) or len(nom.encode()) >= 64:
        raisons.append(f"MONGO_DB « {nom} » : nom de base invalide (ni espace ni "
                       "/ \\ . \" $ * < > : | ?, moins de 64 caractères)")
    if raisons:
        return {"mongo / secrets": raisons}
    uri = os.environ["MONGO_URI"]
    if uri.startswith("mongomock://"):
        # Base en memoire, videe a chaque redemarrage : la demo, serve et les
        # tests la declarent (BASE_EN_MEMOIRE=1). Recopiee par erreur dans le
        # .env d'un client, elle perdrait ses dossiers sans un mot.
        if os.environ.get("BASE_EN_MEMOIRE") != "1":
            return {"mongo / secrets": ["MONGO_URI=mongomock:// est une base en mémoire, "
                                        "vidée à chaque redémarrage : réservée à la démo "
                                        "et aux tests"]}
    elif not uri.startswith(("mongodb://", "mongodb+srv://")):
        return {"mongo / secrets": ["MONGO_URI doit commencer par mongodb:// ou "
                                    "mongodb+srv://"]}
    # Hote vide : pymongo leverait InvalidURI en pleine verification. Lu a la
    # main, sans MongoClient qui resout le DNS d'une adresse mongodb+srv://.
    elif not uri.split("://", 1)[1].split("/", 1)[0].rsplit("@", 1)[-1].strip(", "):
        return {"mongo / secrets": ["MONGO_URI sans adresse de serveur après mongodb://"]}
    return {}


def _verifier_compression():
    """Clé absente = "couleur"."""
    c = instance().get("compression", "couleur")
    if c not in ("couleur", "gris"):
        return {"compression": [f"attend \"couleur\" ou \"gris\" (reçu : {c!r})"]}
    return {}


def _verifier_conservation():
    """Conservation des donnees : au bout de combien de jours un dossier termine
    voit ses pieces effacees. La FORME avant les
    valeurs, sinon c'est le garde-fou qui plante sur une config malformee.

    Bloc absent = aucune purge : les instances installees avant ce bloc
    demarrent sans rien changer."""
    c = instance().get("conservation")
    if c is None:
        return {}
    if not isinstance(c, dict):
        return {"conservation": ["attend un objet, ex. {\"jours\": 1095, "
                                 "\"jours_candidature\": 730, \"apres\": "
                                 "[\"RemisComptable\", \"Rejetee\", \"Abandonnee\"]}"]}
    import store          # store importe config : import local, cycle sinon
    raisons = []
    for cle in ("jours", "jours_candidature"):
        v = c.get(cle)
        # isinstance(True, int) est vrai -- refuser bool explicitement.
        if not isinstance(v, int) or isinstance(v, bool) or v <= 0:
            raisons.append(f"« {cle} » attend un entier > 0 (reçu : {v!r})")
    apres = c.get("apres")
    if (not isinstance(apres, list) or not apres
            or not all(isinstance(e, str) for e in apres)):
        raisons.append("« apres » attend une liste non vide d'états")
    elif inconnus := [e for e in apres if e not in store.ETATS]:
        raisons.append("« apres » : état(s) inconnu(s) " + ", ".join(inconnus)
                       + " — au choix : " + ", ".join(store.ETATS))
    return {"conservation": raisons} if raisons else {}


def _verifier_sharepoint():
    """Bouton « Récupérer les dossiers ». Patron de _verifier_conservation.
    Bloc absent = pas de bouton."""
    c = instance().get("sharepoint")
    if c is None:
        return {}
    if not isinstance(c, dict):
        return {"sharepoint": ["attend un objet, ex. {\"site\": "
                               "\"https://acme.sharepoint.com/sites/rh\"}"]}
    raisons = []
    site = c.get("site")
    if not isinstance(site, str) or not re.match(r"https://[^/]+\.sharepoint\.com/", site):
        raisons.append(f"« site » attend l'adresse du site, ex. "
                       f"https://acme.sharepoint.com/sites/rh (reçu : {site!r})")
    if not (isinstance(c.get("racine"), str) and c["racine"].strip("/ ")):
        raisons.append("« racine » attend le dossier qui contient les groupes "
                       "d'établissements, ex. \"DOSSIERS SALARIES\"")
    if "client_id" in c and not (isinstance(c["client_id"], str) and c["client_id"]):
        raisons.append("« client_id » attend une chaîne (ou rien : client Office)")
    for cle, ex in (("ignorer", "[\"Anciens\", \"Archives\"]"), ("mots_ignores", "[\"restaurant\"]")):
        v = c.get(cle, [])
        if not (isinstance(v, list) and all(isinstance(x, str) for x in v)):
            raisons.append(f"« {cle} » attend une liste de textes, ex. {ex}")
    postes = c.get("postes", {})
    if not (isinstance(postes, dict)
            and all(isinstance(k, str) and isinstance(v, str) for k, v in postes.items())):
        raisons.append("« postes » attend {\"dossier SharePoint\": \"poste du formulaire\"}, "
                       "ex. {\"Chefs d'équipe\": \"Manager\", \"Direction\": \"\"}")
    dossiers = c.get("dossiers", {})
    if not (isinstance(dossiers, dict)
            and all(isinstance(k, str) and isinstance(v, str) for k, v in dossiers.items())):
        raisons.append("« dossiers » attend {\"Société / Établissement\": \"GROUPE/Dossier\"}, "
                       "ex. {\"ACME Sud / Marseille Prado\": \"RESTAURANTS/Marseille\"}")
    return {"sharepoint": raisons} if raisons else {}


# Les modeles que l'app envoie a chaque dossier (app.py : _mail). Un modele
# absent ne casse rien au demarrage : le mail echoue plus tard, note au journal
# du dossier, et personne n'est prevenu (config_wingstop sans demande_validee).
# recap_hebdo n'en est pas : il part d'un cron (recap.py), son echec s'affiche
# deja en bandeau sur /salaries ; configurer.avertissements() le signale.
MODELES_MAIL = ("nouvelle_soumission", "demande_validee", "rejet", "rappel_dpae")


def modele_mail_manquant(m):
    """config/mails/<m>.txt -> "absent", "vide", ou None s'il est utilisable."""
    p = CLIENT / "mails" / f"{m}.txt"
    if not p.is_file():
        return "absent"
    return None if p.read_text(encoding="utf-8").strip() else "vide"


def _verifier_mails():
    manques = [f"mails/{m}.txt {q}" for m in MODELES_MAIL if (q := modele_mail_manquant(m))]
    return {"modèles de mail": manques} if manques else {}


def _verifier_theme():
    """La DA du client. Patron de _verifier_conservation : la FORME avant les
    valeurs, et un KO propre qui nomme le sujet.

    Refus au demarrage plutot qu'ignore en silence -- une couleur ignoree,
    c'est un client qui ouvre son app le premier jour, la voit verte, et
    personne ne sait pourquoi ; un logo absent, c'est un en-tete casse devant
    des candidats. Fichier absent = la DA d'origine, les instances
    installees avant theme.json demarrent sans rien changer.

    Les regles vivent dans apparence.py, ici on ne fait que les rapporter."""
    try:
        raisons = resoudre(theme(), MARQUE).raisons
    except ValueError as e:                     # JSON casse : theme() le laisse passer
        raisons = [f"illisible — ce n'est pas du JSON valide : {e}"]
    return {"theme.json": raisons} if raisons else {}



def _verifier_installation():
    manques = {}
    # Sans destinataire RH, chaque soumission part en `mail_echoue` dans le
    # journal et personne ne le voit : deux instances de test ont servi comme
    # ca, verifier() disant OK. Un mail sans expediteur est refuse par un
    # relais SMTP reel, autant le dire avant le passage en mode smtp.
    m = instance().get("mails") or {}
    if not [d for d in (m.get("rh") or []) if d]:
        manques["mails"] = ["« rh » vide : aucune adresse ne recevrait les demandes "
                            "à valider — renseignez mails.rh dans instance.json"]
    if not (m.get("expediteur") or "").strip():
        manques.setdefault("mails", []).append(
            "« expediteur » vide : renseignez l'adresse d'envoi dans mails")
    if (m.get("mode") == "smtp" and m.get("utilisateur")
            and not os.environ.get("SMTP_MOT_DE_PASSE")):
        manques["SMTP_MOT_DE_PASSE"] = ["absent de l'environnement (relais authentifié)"]
    if ((instance().get("signature") or {}).get("mode") == "yousign"
            and not os.environ.get("YOUSIGN_CLE_API")):
        manques["signature"] = ["YOUSIGN_CLE_API absent de l'environnement"]
    # Le mail hebdo au cabinet est fabrique hors requete (cron) : sans base
    # d'URL il partirait avec des liens vides, et personne ne lit la sortie
    # d'un cron -- meme trou que mails.rh, refuse au meme endroit.
    if not url_publique():
        manques["url"] = ["vide : l'adresse publique de l'app (ex. https://embauche.acme.fr), "
                          "base des liens de tous les mails"]
    return manques


def _base_illisible(e):
    """Manque commun quand la base leve en pleine verification : URI que
    pymongo refuse (port non numerique, options mal placees...) ou serveur
    qui ne repond pas. Le type seul : le message peut citer l'adresse."""
    return {"mongo / secrets": [f"base mongo illisible ({type(e).__name__}) : "
                                "vérifiez MONGO_URI et MONGO_DB, puis que la base répond"]}


def _verifier_societes():
    """Les societes en base : au moins un etablissement, chacun nomme (le
    formulaire et _verifier_regles en dependent). Sans base configuree : un
    manque, pas une trace -- comme _verifier_comptes."""
    try:
        data = societes()
    except RuntimeError as e:
        return {"établissements": [str(e)]}
    except (PyMongoError, ValueError) as e:     # ValueError : port non numerique
        return _base_illisible(e)
    raisons = []
    for i, s in enumerate(data, 1):
        nom = s.get("nom")
        etabs = s.get("etablissements")
        if not isinstance(nom, str) or not nom.strip():
            raisons.append(f"société n° {i} sans nom")
        elif not isinstance(etabs, list) or any(
                not isinstance(e, dict) or not isinstance(e.get("nom"), str)
                or not e["nom"].strip() for e in etabs):
            raisons.append(f"« {nom} » : un établissement sans nom ou mal formé")
        else:
            # config.mentions() fusionne ces dicts : un null ou une liste
            # ferait planter la generation du contrat APRES la validation.
            for quoi, x in [(nom, s)] + [(f"{nom} / {e['nom']}", e) for e in etabs]:
                if "mentions" in x and not isinstance(x["mentions"], dict):
                    raisons.append(f"« {quoi} » : mentions doit être un objet "
                                   "{clé: valeur}")
    if not raisons and not any(s["etablissements"] for s in data):
        raisons.append("aucun établissement en base — lancez : python configurer.py "
                       "--assister (ou importer-societes <fichier.json>)")
    return {"établissements": raisons} if raisons else {}


def _verifier_comptes():
    """Aucun compte en base = refus au demarrage ; un compte sans hash valide
    aussi, plutot qu'un login() qui plante. Sans base configuree : un manque,
    pas une trace -- doctor et configurer doivent rendre leur bilan."""
    from werkzeug.security import check_password_hash
    try:
        d = comptes()
    except RuntimeError as e:
        return {"comptes": [str(e)]}
    except (PyMongoError, ValueError) as e:
        return _base_illisible(e)
    if not d:
        return {"comptes": ["aucun compte de connexion en base — lancez : "
                            "python configurer.py compte rh"]}
    # Un hash werkzeug a la forme « methode$sel$empreinte » ; une autre forme,
    # ou une methode inconnue (werkzeug leve ValueError), et personne ne peut
    # plus se connecter, sans message.
    raisons, par_defaut = [], False
    for ident, u in d.items():
        h = u.get("mdp_hash")
        try:
            if not isinstance(h, str) or h.count("$") < 2:
                raise ValueError(h)
            par_defaut = check_password_hash(h, MDP_PAR_DEFAUT) or par_defaut
        except ValueError:
            raisons.append(f"« {ident} » : mdp_hash absent ou invalide")
    if raisons:
        return {"comptes": raisons}
    if par_defaut:
        return {"compte RH": ["mot de passe par défaut — changez-le : "
                              "python configurer.py compte <identifiant>"]}
    return {}


def _verifier_derives():
    """Verifie la FORME des regles derives avant leurs valeurs -- sinon c'est le
    garde-fou lui-meme qui plante sur `si: ["nom"]`."""
    import contrat
    manques = {}
    for d in derives():
        ph = d.get("placeholder")
        raisons = [] if ph else ["règle sans « placeholder » : rien à alimenter"]
        if "format" in d:
            if d["format"] not in contrat.FORMATEURS:
                raisons.append(f"formateur « {d['format']} » inconnu — au choix : "
                               + ", ".join(sorted(contrat.FORMATEURS)))
            if not d.get("de"):
                raisons.append("un « format » exige « de » : le champ à formater")
        else:
            if "alors" not in d:
                raisons.append("ni « format » ni « alors » : la règle ne produit rien")
            si = d.get("si")
            if si is None:
                pass
            elif not isinstance(si, (list, tuple)) or len(si) != 3:
                raisons.append("« si » attend [champ, opérateur, valeur]")
            elif si[1] not in contrat.OPERATEURS:
                raisons.append(f"opérateur « {si[1]} » inconnu — au choix : "
                               + ", ".join(contrat.OPERATEURS))
        if raisons:
            manques[f"derives → {ph or '?'}"] = raisons
    return manques


def _verifier_ids_champs():
    """store ecrit `champs.<id>` : un id avec « . » creerait un sous-document, un
    « $ » en tete serait lu comme un operateur Mongo."""
    mauvais = [c["id"] for c in champs()
               if not isinstance(c.get("id"), str) or "." in c["id"] or c["id"].startswith("$")]
    return {"champs": [f"id « {i} » : ni « . » ni « $ » en tête dans un id de champ"
                       for i in mauvais]} if mauvais else {}


def _verifier_roles():
    """Les roles sont la promesse « n'importe quelle PME » : le moteur ne lit que
    des roles, jamais un id de champ. La FORME avant les valeurs."""
    manques = {}
    declares = formulaire().get("roles", {})
    if not isinstance(declares, dict):
        manques["roles"] = ["formulaire.json → « roles » attend un objet "
                            "{rôle: id de champ}"]
        return manques
    ids = {c["id"] for c in champs()}
    raisons = []
    for nom, cid in declares.items():
        if nom not in ROLES:
            raisons.append(f"rôle « {nom} » inconnu — au choix : "
                           + ", ".join(sorted(ROLES)))
        elif not isinstance(cid, str) or not cid:
            raisons.append(f"rôle « {nom} » : « {cid} » n'est pas un id de champ")
        elif cid not in ids:
            raisons.append(f"rôle « {nom} » → « {cid} », qui n'est pas un "
                           "champ du formulaire")
    for nom in ROLES_REQUIS:
        if role(nom) not in ids:
            raisons.append(f"rôle « {nom} » : aucun champ « {role(nom)} » — "
                           f"déclarez-le dans formulaire.json → roles")
    if raisons:
        manques["roles"] = raisons
    return manques


def _jetons_manquants(chemin, sources):
    """Jetons de `chemin` qu'aucune `sources` n'alimente. None si le fichier
    est absent (appelant : le signaler autrement), [] si tout est couvert."""
    if not chemin.exists():
        return None
    raisons = []
    for j in sorted(_jetons(chemin)):
        if j in sources:
            continue
        proche = next((s for s in sources if s.casefold() == j.strip().casefold()),
                      None)
        raisons.append(f"{{{{{j}}}}} → écrire exactement {{{{{proche}}}}}"
                       if proche else f"{{{{{j}}}}} — aucune source ne l'alimente")
    return raisons


def _verifier_placeholders():
    """Un placeholder de template actif qu'aucune source n'alimente : la
    promesse d'adaptabilite du produit, verifiee au demarrage. La fiche
    salarie du CLIENT en est ; pas le modele generique, qui retire a la
    generation les lignes que la config n'alimente pas (contrat.elaguer)."""
    manques = {}
    sources = placeholders_connus()
    for nom in _templates_actifs():
        raisons = _jetons_manquants(CLIENT / "contrats" / nom, sources)
        if raisons is None:
            manques[nom] = ["fichier absent de config/contrats/"]
        elif raisons:
            manques[nom] = raisons
    if nom_fiche := instance().get("fiche_salarie"):
        raisons = _jetons_manquants(CLIENT / "contrats" / nom_fiche, sources)
        if raisons is None:
            manques[nom_fiche] = ["fichier absent de config/contrats/"]
        elif raisons:
            manques[nom_fiche] = raisons
    elif not MODELE_FICHE_GENERIQUE.exists():
        # Rien a verifier sur ses jetons (elaguer retire a la generation les
        # lignes que la config n'alimente pas) -- seulement sa PRESENCE : un
        # deploiement sans le dossier « modeles/ » demarrait vert, puis ratait
        # la fiche a chaque validation, devant la RH.
        manques["fiche salarié"] = [
            f"fichier absent : {MODELE_FICHE_GENERIQUE} — copiez le dossier "
            "« modeles/ » à côté du code"]
    return manques


def _verifier_regles():
    """Regles `templates` (forme liste) : une valeur de « quand » qui ne peut
    JAMAIS correspondre a un dossier reel passait ici sans message -- le
    mauvais modele partait en silence sur la regle fourre-tout suivante. Cas
    reel : la cle d'un etablissement est « Societe / Etablissement »
    (etablissements()), pas le nom du site seul. La FORME avant les valeurs,
    comme les autres _verifier_* : une regle ou un « quand » mal forme ne doit
    pas faire planter le garde-fou lui-meme."""
    n = _normaliser
    t = instance().get("templates")
    if not isinstance(t, list):
        return {}
    ids = {c["id"] for c in champs()}
    cles_etab = {cle for cle, _ in etablissements()}
    # nom du site seul (sans la societe) -> cle complete, pour le message d'aide
    par_nom_seul = {}
    for s in societes():
        for e in s["etablissements"]:
            par_nom_seul.setdefault(n(e["nom"]), f"{s['nom']} / {e['nom']}")

    manques = {}
    for i, regle in enumerate(t):
        sujet = f"templates → règle {i + 1}"
        if not isinstance(regle, dict):
            manques[sujet] = ["attend un objet {\"quand\": {...}, \"modele\": ...}"]
            continue
        quand = regle.get("quand", {})
        if not isinstance(quand, dict):
            manques[sujet] = ["« quand » attend un objet {champ: valeur}"]
            continue
        raisons = []
        for cid, val in quand.items():
            if cid not in ids:
                raisons.append(f"« {cid} » n'est pas un champ du formulaire")
                continue
            c = champ(cid)
            if c.get("type") == "etablissement":
                if any(n(val) == n(cle) for cle in cles_etab):
                    continue
                if n(val) in par_nom_seul:
                    raisons.append(f"« {cid} » : « {val} » ne correspond à aucun "
                                   f"établissement — écrire exactement "
                                   f"« {par_nom_seul[n(val)]} »")
                else:
                    raisons.append(f"« {cid} » : « {val} » n'est pas une clé "
                                   "d'établissement connue (sociétés en base)")
            elif c.get("options"):
                if not any(n(val) == n(o) for o in c["options"]):
                    raisons.append(f"« {cid} » : « {val} » ne correspond à aucune "
                                   "option (" + ", ".join(map(str, c["options"])) + ")")
        if raisons:
            manques[sujet] = raisons
    return manques


def _verifier_postes():
    """Chaque poste du formulaire doit atteindre un modele."""
    manques = {}
    t = regles()
    vises = set()
    fourre_tout = any(not r.get("quand") for r in t)
    for r in t:
        vises |= {_normaliser(v) for v in r.get("quand", {}).values()}
    for c in champs():
        if c["id"] != role("poste"):
            continue
        for opt in c.get("options", []):
            if not fourre_tout and _normaliser(opt) not in vises:
                manques[f"poste « {opt} »"] = ["aucune règle de template ne le vise"]
    return manques


def _verifier_grille():
    """Chaque poste du formulaire doit avoir SA ligne dans grille.json.

    La grille reconnait un poste par mots normalises, la plus specifique
    gagne -- souple exprès (« Equipier Polyvalent » prend le tarif
    « Equipier »), mais un poste PRIVE de sa ligne tombe alors en silence sur
    celle d'un poste dont l'intitule est contenu dans le sien : « Apprenti
    couvreur » sort au tarif « Couvreur ». Le contrat n'est pas vide, il est
    FAUX -- invisible pour doctor comme pour la RH, et sur le seul champ dont
    personne ne doute. Deux postes qui visent la meme ligne sont donc refuses
    ici : si le partage est voulu, deux lignes de meme montant le disent.
    """
    g = grille()
    if not g:
        return {}
    import contrat
    # La FORME avant les valeurs : sans ca, une ligne sans « poste » fait
    # planter le garde-fou (KeyError) au lieu d'etre refusee -- et le meme
    # acces non garde est dans contrat.ligne_grille().
    lignes = g.get("postes")
    if not isinstance(lignes, list) or not all(
            isinstance(e, dict) and isinstance(e.get("poste"), str) for e in lignes):
        return {"grille.json": ["« postes » attend une liste d'objets portant "
                                "chacun un « poste » (texte)"]}
    manques, pris = {}, {}
    for c in champs():
        if c["id"] != role("poste"):
            continue
        for opt in c.get("options", []):
            ligne = contrat.ligne_grille(opt, g)
            if ligne is None:
                manques[f"grille — poste « {opt} »"] = [
                    "aucune ligne de config/grille.json ne le rémunère"]
            elif (autre := pris.get(ligne["poste"])) is not None:
                # On ne sait pas lequel des deux est le proprietaire legitime de
                # la ligne (ca depend de l'ordre des options) : nommer les deux
                # plutot que d'accuser au hasard celui vu en second.
                manques[f"grille — postes « {autre} » et « {opt} »"] = [
                    f"partagent la ligne « {ligne['poste']} » de config/grille.json "
                    f"— l'un des deux doit avoir sa propre ligne (même montant si "
                    f"le partage est voulu)"]
            else:
                pris[ligne["poste"]] = opt
    for ligne in lignes:
        # « mensuel » l'emporte sur « bareme » dans contrat.salaire()
        if "bareme" in ligne and "mensuel" not in ligne and ligne["poste"] in pris:
            manques.update(_verifier_bareme(ligne, g))
    return manques


def _verifier_bareme(ligne, g):
    """Chaque duree hebdo que le formulaire PROPOSE doit avoir sa ligne de
    bareme. doctor n'essaie que la premiere option : sur config.demo, 7 durees
    sur 9 etaient hors bareme, bilan OK, et le premier salarie reel a 20H avait
    son contrat refuse devant la RH (salaire vide). La FORME avant les valeurs.

    « lettres » n'est plus une cle attendue -- contrat.salaire() la derive de
    « chiffres » -- mais reste toleree (ignoree) si un client l'a encore.

    simplification volontaire : une duree en saisie libre (champ sans options)
    n'est pas enumerable, donc pas verifiee ici -- doctor la couvre pour sa
    valeur d'exemple."""
    import contrat
    sujet = f"grille — barème « {ligne['poste']} »"
    bareme, cid = ligne["bareme"], g.get("champ_heures")
    if not isinstance(bareme, dict) or not all(
            isinstance(v, dict) and isinstance(v.get("chiffres"), str) and v["chiffres"]
            for v in bareme.values()):
        return {sujet: ["« bareme » attend un objet {\"<heures>\": {\"chiffres\": \"…\"}}, "
                        "chaque ligne avec son montant"]}
    # La FORME ne suffit pas : un « chiffres » non vide mais illisible tombait a
    # 0 dans contrat._nombre SANS ERREUR, et le contrat sortait signe avec
    # « (zéro) » en toutes lettres a cote du bon montant en chiffres.
    def lisible(s):
        # Chiffres et separateurs SEULEMENT : _nombre() extrait « 1500 » de
        # « a partir de 1500 », ce qui mettrait de la prose dans la case montant
        # du contrat a cote d'un « mille cinq cents » sorti de nulle part.
        return re.fullmatch(r"[\d ., ]+", str(s)) and contrat._nombre(s)

    if illisibles := sorted(h for h, v in bareme.items() if not lisible(v["chiffres"])):
        return {sujet: [f"montant illisible pour « {h} » : {bareme[h]['chiffres']!r} — "
                        "attendu un nombre (ex. « 1 280,24 » ou « 1.280,24 »), sinon le "
                        "contrat porterait « zéro » en toutes lettres"
                        for h in illisibles]}
    if not isinstance(cid, str) or cid not in {c["id"] for c in champs()}:
        return {sujet: [f"« champ_heures » ({cid!r}) doit nommer le champ du formulaire "
                        "qui porte la durée hebdomadaire"]}
    trous = [o for o in champ(cid).get("options") or []
             if contrat.cle_bareme(o) not in bareme]
    if trous:
        return {sujet: [f"durée(s) proposée(s) par le formulaire sans ligne : "
                        f"{', '.join(map(str, trous))} — ces salariés auraient un "
                        f"contrat sans salaire"]}
    return {}


def verifier():
    """Garde-fou de demarrage. Renvoie {} si tout va bien, sinon un dict
    {sujet: [raisons]} et l'instance ne sert pas."""
    manques = _verifier_mongo()
    if not manques:
        manques.update(_verifier_societes())
    # Base absente ou illisible : comptes et societes rediraient le meme
    # manque, et les checks qui lisent les societes leveraient.
    sans_base = "mongo / secrets" in manques
    lisent_societes = {_verifier_placeholders, _verifier_regles}
    lisent_base = lisent_societes | {_verifier_comptes}
    for check in (_verifier_version, _verifier_comptes,
                  _verifier_conservation, _verifier_compression, _verifier_sharepoint, _verifier_theme,
                  _verifier_mails,
                  _verifier_installation,
                  _verifier_derives, _verifier_ids_champs, _verifier_roles, _verifier_placeholders,
                  _verifier_regles, _verifier_postes, _verifier_grille):
        # Societes absentes ou mal formees : ces checks leveraient, et le
        # manque « établissements » dit deja quoi faire.
        if (sans_base and check in lisent_base) or (
                "établissements" in manques and check in lisent_societes):
            continue
        manques.update(check())
    return manques
