"""`doctor` : vérifier la couverture d'une config en PRODUISANT les contrats.

    python tests/test_doctor.py

Le livrable de la relecture juridique : le client relit les contrats que sa
config produit vraiment, pas un modèle abstrait. Ce que le test prouve :

  - chaque croisement poste × établissement produit un fichier, ou dit pourquoi
    il n'en produit pas (aucun modèle visé, valeur critique vide) ;
  - les axes du tableau sont DÉRIVÉS des règles de template du client, jamais
    codés en dur — un client qui branche sur le temps partiel voit ses deux cas,
    un client qui branche sur l'établissement voit les siens, et un poste en
    texte libre reste couvert par les postes que ses règles nomment ;
  - un établissement en contrat déposé est ignoré, pas compté en échec ;
  - doctor ne touche pas à data/ : il tourne chez un client en production.
"""
import json
import os
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
BASE = Path(tempfile.mkdtemp(prefix="contratgen-doctor-"))
os.environ["CONFIG_DIR"] = str(BASE / "config")
os.environ["DONNEES"] = str(BASE / "data")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from werkzeug.security import generate_password_hash          # noqa: E402
import config                                                 # noqa: E402
import contrat                                                # noqa: E402
import doctor                                                 # noqa: E402
import store                                                  # noqa: E402

CONF = BASE / "config"
(CONF / "contrats").mkdir(parents=True)

INSTANCE = {
    "config_version": 1,
    "client": "ACME",
    "secret": "a" * 64, "url": "http://localhost",
    "signature": {"mode": "manuel"},
    "mails": {"mode": "console", "hote": "", "port": 587, "utilisateur": "",
              "mot_de_passe": "", "expediteur": "rh@acme.example", "rh": ["rh@acme.example"],
              "dpae": None, "recap": None, "comptable_defaut": "c@acme.example"},
    "utilisateurs": {"rh": {"mdp_hash": generate_password_hash("pas-demo")}},
    "motifs_ko": ["Autre"],
    "templates": {"Manager": "contrat.txt"},
}
FORMULAIRE = {"champs": [
    {"id": "etablissement", "libelle": "Établissement", "type": "etablissement",
     "requis": True},
    {"id": "nom", "libelle": "Nom", "type": "texte", "requis": True,
     "placeholder": "Nom"},
    {"id": "poste", "libelle": "Poste", "type": "choix", "requis": True,
     "options": ["Manager", "Leavers"], "placeholder": "Poste"},
    {"id": "temps_partiel", "libelle": "Temps partiel", "type": "choix",
     "options": ["OUI", "NON"]},
]}
SOCIETES = [{"nom": "ACME", "siren": "111 111 111",
             "comptable_email": "c@acme.example",
             "mentions": {"RaisonSociale": "ACME SAS"},
             "etablissements": [
                 {"nom": "Siège", "siret": "111 111 111 00011",
                  "mentions": {"AdresseEtablissement": "1 rue X"}},
                 {"nom": "Entrepôt", "siret": "111 111 111 00029", "contrat": "depose",
                  "mentions": {"AdresseEtablissement": "2 rue Y"}}]}]

MODELE = "Bonjour {{Nom}}, poste {{Poste}} chez {{RaisonSociale}}."


def ecrire(instance=None, modele=MODELE, grille=None, formulaire=None, societes=None):
    (CONF / "instance.json").write_text(
        json.dumps({**INSTANCE, **(instance or {})}, ensure_ascii=False),
        encoding="utf-8")
    (CONF / "formulaire.json").write_text(
        json.dumps(formulaire or FORMULAIRE), encoding="utf-8")
    (CONF / "societes.json").write_text(
        json.dumps(societes or SOCIETES), encoding="utf-8")
    (CONF / "contrats" / "contrat.txt").write_text(modele, encoding="utf-8")
    f = CONF / "grille.json"
    if grille is None:
        f.unlink(missing_ok=True)
    else:
        f.write_text(json.dumps(grille), encoding="utf-8")
    config._CACHE.clear()


def test_chaque_cas_produit_un_contrat_relisible():
    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"}})
    lignes = doctor.examiner()

    from docx import Document
    import contrat as moteur

    ok = [l for l in lignes if l["statut"] == "ok"]
    assert len(ok) == 2, [l["libelle"] for l in lignes]      # 2 postes × 1 étab généré
    for l in ok:
        assert Path(l["fichier"]).exists(), l
        texte = "\n".join(p.text for p in moteur.paragraphes(Document(l["fichier"])))
        assert "ACME SAS" in texte


def test_etablissement_en_contrat_depose_ignore_pas_en_echec():
    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"}})
    lignes = doctor.examiner()
    ignores = [l for l in lignes if l["statut"] == "ignore"]
    assert ignores and all("Entrepôt" in l["libelle"] for l in ignores), lignes
    assert doctor.verdict(lignes) == 0, "un couloir « contrat déposé » compte en échec"


def test_poste_sans_modele_signale_avant_la_mise_en_service():
    ecrire()                                     # « Leavers » absent des templates
    lignes = doctor.examiner()
    l = next(l for l in lignes if l["champs"]["poste"] == "Leavers")
    assert l["statut"] == "sans_modele", l
    assert doctor.verdict(lignes) != 0, "un poste sans modèle passe en vert"


def test_valeur_critique_vide_refuse_le_contrat():
    """Le mode d'échec de wingstop_ : un contrat qui sort avec une ligne de paie
    blanche. Poste absent de la grille -> pas de salaire -> refus, pas un .docx
    silencieusement troué."""
    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"},
            "critiques": ["SalaireChiffres"]},
           modele=MODELE + " Rémunération {{SalaireChiffres}} €.",
           grille={"postes": [{"poste": "Directeur", "mensuel": 3000}]})
    l = next(l for l in doctor.examiner() if l["champs"]["poste"] == "Manager")
    assert l["statut"] == "refuse", l
    assert "SalaireChiffres" in l["detail"], l


def test_axes_derives_des_regles_du_client():
    """Un client qui branche sur le temps partiel doit voir SES deux cas : les
    axes du tableau sortent de `quand`, ils ne sont pas codés dans le produit."""
    ecrire({"templates": [
        {"quand": {"poste": "Manager", "temps_partiel": "OUI"}, "modele": "contrat.txt"},
        {"quand": {"poste": "Manager"}, "modele": "contrat.txt"},
        {"quand": {"poste": "Leavers"}, "modele": "contrat.txt"}]})
    lignes = [l for l in doctor.examiner() if l["statut"] != "ignore"]
    managers = [l for l in lignes if l["champs"]["poste"] == "Manager"]
    assert {l["champs"]["temps_partiel"] for l in managers} == {"OUI", "NON"}, managers
    assert len(lignes) == 4, [l["libelle"] for l in lignes]   # 2 postes × 2 temps


def test_regle_visant_un_etablissement_est_couverte():
    """`modele_pour` accepte une règle qui vise l'établissement (« un modèle par
    établissement quand les mentions sont figées dans la prose », config.py).
    L'établissement est déjà l'axe extérieur : le remettre en colonne écrasait
    la vraie clé par une valeur d'exemple vide, et AUCUNE règle ne matchait —
    trou réel maquillé en ✗, ou masqué par un fourre-tout."""
    trois = [{**SOCIETES[0], "etablissements": SOCIETES[0]["etablissements"] + [
        {"nom": "Atelier", "siret": "111 111 111 00037",
         "mentions": {"AdresseEtablissement": "3 rue Z"}}]}]
    ecrire({"templates": [{"quand": {"etablissement": "ACME / Siège"},
                           "modele": "contrat.txt"}]}, societes=trois)
    par_etab = {l["libelle"].split(" / ", 1)[-1]: l["statut"]
                for l in doctor.examiner()}
    assert par_etab["ACME / Siège"] == "ok", par_etab
    assert par_etab["ACME / Atelier"] == "sans_modele", par_etab
    assert par_etab["ACME / Entrepôt"] == "ignore", par_etab


def test_poste_hors_options_reste_couvert():
    """Un client dont le poste est du texte libre (pas un `choix` à options) :
    les postes à couvrir sont ceux que ses règles nomment, sinon doctor teste un
    seul poste inventé et annonce « tout est couvert » sur un cas fictif."""
    libre = {"champs": [c if c["id"] != "poste" else
                        {"id": "poste", "libelle": "Poste", "type": "texte",
                         "requis": True, "placeholder": "Poste"}
                        for c in FORMULAIRE["champs"]]}
    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"}},
           formulaire=libre)
    postes = {l["champs"]["poste"] for l in doctor.examiner()}
    assert postes == {"Manager", "Leavers"}, postes


def test_rapport_ne_compte_pas_les_hors_perimetre():
    """`rapport()` n'etait appele par aucun test — seuls les statuts d'examiner()
    l'etaient. Resultat : le resume annoncait « 3 cas sur 9 produisent un
    contrat » alors que ces 3 etaient justement les etablissements en contrat
    depose, qui n'en produisent aucun et n'ont pas a en produire."""
    ecrire()                                     # Manager ✓, Leavers ✗, 2 déposés
    texte = doctor.rapport(doctor.examiner())
    assert "1 cas sur 2 produisent un contrat (2 en contrat déposé, hors compte)."         in texte, texte

    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"}})
    texte = doctor.rapport(doctor.examiner())
    assert "Les 2 cas produisent un contrat (2 en contrat déposé, hors compte)."         in texte, texte


def test_rapport_quand_aucun_contrat_n_est_a_produire():
    """Une PME qui utilise l'app pour collecter et suivre, mais fait TOUS ses
    contrats à la main : `attendus` vaut 0 et le résumé annonçait « Les 0 cas
    produisent un contrat », une affirmation sur un ensemble vide."""
    tout_depose = [{**SOCIETES[0], "etablissements": [
        {**e, "contrat": "depose"} for e in SOCIETES[0]["etablissements"]]}]
    ecrire(societes=tout_depose)
    lignes = doctor.examiner()
    assert {l["statut"] for l in lignes} == {"ignore"}, lignes
    assert doctor.verdict(lignes) == 0, "aucun contrat attendu ne peut pas être un échec"
    assert "Aucun contrat à produire (4 en contrat déposé, hors compte)." \
        in doctor.rapport(lignes), doctor.rapport(lignes)


def test_regle_exclusion_volontaire_hors_compte():
    """Un cas volontairement sans modèle (« Manager » n'existe qu'à temps plein
    chez ce client, pas de contrat « Manager » à temps partiel) ne doit plus
    sortir en ✗ irrécupérable : `modele: null` le marque « exclu (volontaire) »,
    statut à part, hors compte comme un contrat déposé, sans faire échouer
    doctor ni bloquer configurer.py."""
    ecrire({"templates": [
        {"quand": {"poste": "Manager", "temps_partiel": "OUI"}, "modele": None},
        {"quand": {"poste": "Manager", "temps_partiel": "NON"}, "modele": "contrat.txt"},
        {"quand": {"poste": "Leavers"}, "modele": "contrat.txt"}]})
    lignes = doctor.examiner()

    exclus = [l for l in lignes if l["statut"] == "exclu"]
    # 1 seul étab en mode « genere » (Siège) ; l'Entrepôt (déposé) est ignoré
    # avant même d'atteindre la règle.
    assert len(exclus) == 1, lignes
    assert all(l["champs"]["poste"] == "Manager" and l["champs"]["temps_partiel"] == "OUI"
               for l in exclus), exclus
    assert all(l["detail"] == "exclu (volontaire)" for l in exclus), exclus
    assert doctor.verdict(lignes) == 0, "une exclusion volontaire fait échouer doctor"
    # Le garde-fou de demarrage specifique aux postes ne doit pas re-signaler
    # « Manager » comme poste sans regle : la fixture globale echoue par
    # ailleurs (roles nom/email non declares, sans rapport avec ce test).
    assert config._verifier_postes() == {}, config._verifier_postes()

    texte = doctor.rapport(lignes)
    assert "– exclu (volontaire)" in texte, texte
    assert "exclu(s) volontairement" in texte, texte    # hors compte, comme "contrat déposé"


def test_regle_exclusion_volontaire_distincte_d_une_regle_absente():
    """`config.regle_pour` doit distinguer les deux None de `modele_pour` :
    aucune règle ne matche (sans_modele) vs. une règle matche avec `modele:
    null` (exclu). Un vrai dossier refuse dans les deux cas -- app.py ne lit
    que modele_pour, dont le contrat (None) est inchangé."""
    ecrire({"templates": [
        {"quand": {"poste": "Manager", "temps_partiel": "OUI"}, "modele": None},
        {"quand": {"poste": "Manager", "temps_partiel": "NON"}, "modele": "contrat.txt"}]})

    exclu = {"poste": "Manager", "temps_partiel": "OUI"}
    absent = {"poste": "Stagiaire", "temps_partiel": "OUI"}

    regle = config.regle_pour(exclu)
    assert regle is not None and regle["modele"] is None, regle
    assert config.regle_pour(absent) is None

    # Cote generation reelle (ce que lit app.py::_produire_contrat) : refuse
    # pareil dans les deux cas, comme avant l'ajout de l'exclusion.
    assert config.modele_pour(exclu) is None
    assert config.modele_pour(absent) is None


def test_exclusion_volontaire_en_forme_dict_aussi():
    """Les deux formes de `templates` doivent lire `null` pareil : la forme
    liste donnait « exclu (volontaire) », la forme dict « aucun modèle ne vise
    ce cas » -- le relecteur voyait un trou de config là où le client avait
    explicitement dit « ce poste n'existe pas ici »."""
    ecrire({"templates": {"Manager": None, "Leavers": "contrat.txt"}})
    lignes = doctor.examiner()

    exclus = [l for l in lignes if l["statut"] == "exclu"]
    assert len(exclus) == 1, lignes                  # 1 seul étab en mode « genere »
    assert exclus[0]["champs"]["poste"] == "Manager", exclus
    assert exclus[0]["detail"] == "exclu (volontaire)", exclus
    assert not [l for l in lignes if l["statut"] == "sans_modele"], lignes
    assert doctor.verdict(lignes) == 0, "une exclusion volontaire fait échouer doctor"


def test_champ_heures_devient_un_axe():
    """Reproduction du bug réel : une règle branche sur « temps partiel »
    OUI/NON, mais la durée hebdo (qui décide la ligne de barème) n'était pas un
    axe et restait figée sur sa première option -- le cas NON (temps plein)
    sortait un contrat « 151,67 h/mois » payé au tarif temps partiel, doctor
    vert. La durée hebdo doit varier elle aussi."""
    partiel = {"chiffres": "533,00", "lettres": "cinq cent trente-trois euros"}
    plein = {"chiffres": "1 867,00", "lettres": "mille huit cent soixante-sept euros"}
    formulaire = {"champs": FORMULAIRE["champs"] + [
        {"id": "temps_travail", "libelle": "Durée hebdo", "type": "choix",
         "options": ["10H", "35H"]}]}
    ecrire({"templates": [
        {"quand": {"poste": "Manager", "temps_partiel": "OUI"}, "modele": "contrat.txt"},
        {"quand": {"poste": "Manager", "temps_partiel": "NON"}, "modele": "contrat.txt"},
        {"quand": {"poste": "Leavers"}, "modele": "contrat.txt"}]},
           formulaire=formulaire,
           grille={"champ_heures": "temps_travail",
                   "postes": [{"poste": "Manager", "bareme": {"10": partiel, "35": plein}}]})

    assert "temps_travail" in doctor.axes(), doctor.axes()
    lignes = [l for l in doctor.examiner() if l["statut"] != "ignore"]
    manager_non = [l for l in lignes if l["champs"]["poste"] == "Manager"
                   and l["champs"]["temps_partiel"] == "NON"]
    # Les deux durées sont essayées pour CHAQUE branche -- c'est la variante
    # à 35H qui aurait été masquée avant le correctif.
    assert {l["champs"]["temps_travail"] for l in manager_non} == {"10H", "35H"}, manager_non


def test_champs_lus_par_derives_deviennent_des_axes():
    """Les champs que seules des règles `derives` lisent (`si[0]`, `de`) doivent
    varier aussi, sans doublon, et seulement s'ils désignent un vrai champ du
    formulaire -- ni un placeholder calculé, ni l'établissement (déjà l'axe
    extérieur)."""
    formulaire = {"champs": FORMULAIRE["champs"] + [
        {"id": "cdd", "libelle": "CDD ?", "type": "choix", "options": ["OUI", "NON"]}]}
    ecrire({"templates": {"Manager": "contrat.txt"},
            "derives": [
                {"placeholder": "MentionCdd", "si": ["cdd", "==", "OUI"],
                 "alors": "à durée déterminée", "sinon": "à durée indéterminée"},
                {"placeholder": "CddMajuscule", "format": "majuscules", "de": "cdd"},
                # placeholder calculé, pas un champ du formulaire -- ignoré
                {"placeholder": "X", "format": "majuscules", "de": "SalaireChiffres"},
                # l'établissement est déjà l'axe extérieur -- pas un doublon
                {"placeholder": "Y", "format": "majuscules", "de": "etablissement"}]},
           formulaire=formulaire)

    axes = doctor.axes()
    assert axes.count("cdd") == 1, axes                 # cité par « si » ET « de »
    assert "SalaireChiffres" not in axes, axes
    assert "etablissement" not in axes, axes

    valeurs = {l["champs"]["cdd"] for l in doctor.examiner()
               if l["champs"]["poste"] == "Manager"}
    assert valeurs == {"OUI", "NON"}, valeurs


def test_court_ne_tronque_pas_au_milieu_d_un_mot():
    """Un libellé long ne doit plus produire de nom de fichier coupé en plein
    mot (`…-Chamber.docx`) : la coupe tombe sur un séparateur, ou pas de coupe
    du tout."""
    long_libelle = ("Equipier Polyvalent Confirme / OUI / 24H / "
                     "Francais / ACME / Paris Bastille Confluence")
    brut = store.SAIN.sub("-", contrat.sans_accent(long_libelle)).strip("-")
    assert len(brut) > 40, "le fixture doit forcer la troncature"

    court = doctor._court(long_libelle)
    assert len(court) <= 40, court
    assert brut.startswith(court), (brut, court)
    suite = brut[len(court):]
    assert suite == "" or suite.startswith("-"), (brut, court)  # coupe = separateur


def test_fichier_genere_a_un_nom_lisible_non_tronque_en_plein_mot():
    """Bout en bout : un poste au nom long produit un vrai fichier dont le nom
    n'est jamais coupé en plein mot (le poste seul dépasse déjà 40 caractères
    une fois assaini)."""
    poste_long = "Manager Senior Adjoint Confirme Principal"
    formulaire = {"champs": [c if c["id"] != "poste" else
                             {**c, "options": [poste_long]} for c in FORMULAIRE["champs"]]}
    ecrire({"templates": {poste_long: "contrat.txt"}}, formulaire=formulaire)

    ok = [l for l in doctor.examiner() if l["statut"] == "ok"]
    assert ok, ok
    for l in ok:
        stem = Path(l["fichier"]).stem
        court = stem.split("-", 1)[1]                   # retire le préfixe NN-
        brut = store.SAIN.sub("-", contrat.sans_accent(l["libelle"])).strip("-")
        suite = brut[len(court):]
        assert suite == "" or suite.startswith("-"), (brut, court)


def test_fiche_salarie_produite_et_relue():
    """La fiche part à chaque validation : doctor la produit aussi, modèle
    générique par défaut, et une fiche client trouée fait échouer le verdict."""
    from docx import Document
    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"}})
    l = doctor.fiche()
    assert l["statut"] == "ok" and l["detail"] == "fiche_salarie.docx", l
    texte = "\n".join(p.text for p in contrat.paragraphes(Document(l["fichier"])))
    assert "{{" not in texte and "Siège" in texte, texte
    # config minimale : le generique retire ce qu'elle n'alimente pas, et le dit
    assert "Disponibilités" not in texte and "Lundi" in l["retires"], l
    lignes = doctor.examiner() + [l]
    assert doctor.verdict(lignes) == 0
    rapport = doctor.rapport(lignes)
    assert "Les 2 cas produisent un contrat" in rapport
    assert "lignes absentes de la config : " in rapport and "Lundi" in rapport, rapport

    (CONF / "contrats" / "fiche.txt").write_text("Fiche {{Inconnu}}", encoding="utf-8")
    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"},
            "fiche_salarie": "fiche.txt"})
    l = doctor.fiche()
    assert l["statut"] == "refuse", l
    assert doctor.verdict(doctor.examiner() + [l]) != 0, "fiche trouée passe en vert"


def test_fiche_client_voit_la_saisie_rh_comme_les_contrats():
    """`examiner()` injecte les placeholders de `saisie_rh` (ce que la RH tape à
    la génération) ; `fiche()` les oubliait. Une fiche salarié CLIENT qui en
    utilise un passait le garde-fou de démarrage — la source existe — puis
    sortait « refuse » chez doctor : `configurer` annonçait KO sur une config
    saine, et l'installateur cherchait un défaut qui n'existe pas."""
    (CONF / "contrats" / "fiche_rh.txt").write_text(
        "Fiche de {{Nom}} — entrée le {{DateEntree}}", encoding="utf-8")
    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"},
            "saisie_rh": ["DateEntree"], "fiche_salarie": "fiche_rh.txt"})

    assert "fiche_rh.txt" not in config.verifier(), config.verifier()
    ligne = doctor.fiche()
    assert ligne["statut"] == "ok", ligne
    assert doctor.verdict(doctor.examiner() + [ligne]) == 0, ligne


def test_regle_mal_formee_ne_plante_pas_doctor():
    """doctor tourne même sur une config refusée au démarrage : une règle qui
    n'est pas un objet est sautée, pas un AttributeError dans axes()."""
    ecrire({"templates": ["pas un objet", {"quand": "pas un dict", "modele": "contrat.txt"},
                          {"quand": {"poste": "Manager"}, "modele": "contrat.txt"}]})
    assert "templates → règle 1" in config.verifier()
    lignes = doctor.examiner()
    assert any(l["statut"] == "ok" for l in lignes), lignes


def test_modele_absent_marque_refuse_pas_de_crash():
    """Un .docx déclaré mais absent : python-docx lève une erreur hors OSError,
    doctor plantait au lieu de marquer le cas ⚠."""
    ecrire({"templates": {"Manager": "absent.docx", "Leavers": "contrat.txt"},
            "fiche_salarie": "fiche_absente.docx"})
    lignes = doctor.examiner() + [doctor.fiche()]
    refus = [l for l in lignes if l["statut"] == "refuse"]
    assert len(refus) == 2 and all("absent" in l["detail"] for l in refus), lignes
    assert doctor.verdict(lignes) != 0


def test_ne_touche_pas_aux_donnees():
    """doctor tourne chez un client en production, sur des dossiers réels."""
    donnees = BASE / "data"
    (donnees / "soumissions").mkdir(parents=True, exist_ok=True)
    (donnees / "soumissions" / "temoin.json").write_text("{}", encoding="utf-8")
    avant = {p: p.stat().st_mtime_ns for p in donnees.rglob("*")}

    ecrire({"templates": {"Manager": "contrat.txt", "Leavers": "contrat.txt"}})
    doctor.examiner()

    assert {p: p.stat().st_mtime_ns for p in donnees.rglob("*")} == avant, \
        "doctor a écrit dans data/"


if __name__ == "__main__":
    import shutil
    try:
        for nom, fn in sorted(globals().items()):
            if nom.startswith("test_"):
                fn()
                print(f"  ✓ {nom}")
        print("\ndoctor : couverture vérifiée en produisant les contrats, axes "
              "dérivés de la config, data/ intact.")
    finally:
        shutil.rmtree(BASE, ignore_errors=True)
        shutil.rmtree(doctor.DOSSIER, ignore_errors=True)
