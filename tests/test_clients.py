"""Plusieurs clients factices, chacun dans SA copie du dépôt.

    python tests/test_clients.py

Modèle de déploiement : un client = une copie du repo. Ce test le prouve pour
de vrai — pour chaque client factice il fait une copie du code + config/, lance
`installer.py "<nom>"`, écrit une config propre au client (sociétés,
établissements, couloirs), puis joue le parcours complet dans un interpréteur
neuf (tests/_client_parcours.py). Il vérifie enfin que les copies sont
étanches : secrets distincts, dossiers de données distincts, aucune trace du
client précédent.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
RACINE = Path(__file__).resolve().parent.parent
CODE = ["app.py", "config.py", "contrat.py", "mails.py", "signature.py",
        "store.py", "installer.py", "recap.py", "placeholders.py", "doctor.py"]

MENTIONS = ["RaisonSociale", "FormeCapital", "RCS", "SiegeSocial", "ConventionCollective"]


def societe(nom, siren, cabinet, raison, siege, etabs):
    return {"nom": nom, "siren": siren, "comptable_email": cabinet,
            "mentions": {"RaisonSociale": raison, "FormeCapital": "SAS au capital de 10 000 €",
                         "RCS": f"RCS {siege.split(',')[-1].strip()} {siren}",
                         "SiegeSocial": siege,
                         "ConventionCollective": "CCN de la restauration rapide"},
            "etablissements": etabs}


def etab(nom, siret, adresse, mode):
    e = {"nom": nom, "siret": siret, "mentions": {"AdresseEtablissement": adresse}}
    if mode != "genere":
        e["contrat"] = mode
    return e


# Aucun id ci-dessous n'existe dans config.exemple/ : c'est le seul lien du
# moteur avec ce formulaire, la table `roles`.
FORMULAIRE_RENOMME = {
    "version": 1, "titre": "Embauche chantier",
    "intro": "À remplir par le conducteur de travaux.",
    "roles": {"etablissement": "chantier", "poste": "fonction",
              "nom": "identite_nom", "prenom": "identite_prenom",
              "email": "courriel", "date_debut": "demarrage"},
    "champs": [
        {"id": "chantier", "libelle": "Chantier", "type": "etablissement", "requis": True},
        {"id": "courriel", "libelle": "Votre email", "type": "email", "requis": True},
        {"id": "identite_nom", "libelle": "Nom", "type": "texte", "requis": True,
         "placeholder": "NomNaissance"},
        {"id": "identite_prenom", "libelle": "Prénom", "type": "texte", "requis": True,
         "placeholder": "Prenom"},
        {"id": "fonction", "libelle": "Poste occupé", "type": "choix", "requis": True,
         "options": ["Couvreur", "Chef d'équipe"], "placeholder": "Poste"},
        {"id": "demarrage", "libelle": "Date de début", "type": "date", "requis": True,
         "placeholder": "DateDebut"},
        # Champ propre au client : le moteur ne le connait pas, il ne va qu'au
        # template. C'est la moitie « champ libre » du contrat.
        {"id": "numero_carte_btp", "libelle": "N° carte BTP", "type": "texte",
         "placeholder": "CarteBTP"},
        {"id": "piece_identite", "libelle": "Pièce d'identité", "type": "piece_jointe",
         "role": "identite", "requis": True},
        {"id": "releve_bancaire", "libelle": "RIB", "type": "piece_jointe",
         "role": "rib", "requis": True},
    ],
}

CLIENTS = [
    {
        "slug": "kebab",
        "nom": "Le Kebab du Coin",
        "societes": [societe(
            "Kebab du Coin", "111 111 111", "compta@kebab.example",
            "KEBAB DU COIN SAS", "4 rue Centrale, 59100 Roubaix", [
                etab("Roubaix Centre", "111 111 111 00011", "4 rue Centrale, 59100 Roubaix", "genere"),
                etab("Roubaix Gare", "111 111 111 00029", "1 parvis de la Gare, 59100 Roubaix", "depose"),
            ])],
        "attendus": ["KEBAB DU COIN SAS", "111 111 111 00011", "Roubaix Centre"],
        "etab_genere": "Kebab du Coin / Roubaix Centre",
        "etab_depose": "Kebab du Coin / Roubaix Gare",
    },
    {
        "slug": "sushi",
        "nom": "Sushi Express",
        "societes": [
            societe("Sushi Express Paris", "222 222 222", "compta.paris@sushi.example",
                    "SUSHI EXPRESS PARIS SARL", "10 av. de l'Opéra, 75001 Paris", [
                        etab("Opéra", "222 222 222 00013", "10 av. de l'Opéra, 75001 Paris", "genere")]),
            societe("Sushi Express Lyon", "333 333 333", "compta.lyon@sushi.example",
                    "SUSHI EXPRESS LYON SARL", "2 rue de la République, 69002 Lyon", [
                        etab("Bellecour", "333 333 333 00017", "2 rue de la République, 69002 Lyon", "depose")]),
        ],
        "attendus": ["SUSHI EXPRESS PARIS SARL", "222 222 222 00013", "Opéra"],
        "etab_genere": "Sushi Express Paris / Opéra",
        "etab_depose": "Sushi Express Lyon / Bellecour",
    },
    {
        # Le client qui prouve la promesse « n'importe quelle PME » : pas un
        # seul id de champ en commun avec config.exemple/. Seuls les roles
        # relient son formulaire au moteur. Avant, « etablissement » et
        # « poste » etaient lus en dur -> KeyError en pleine action RH.
        "slug": "btp",
        "nom": "Toitures du Nord",
        "formulaire": FORMULAIRE_RENOMME,
        "postes": ["Couvreur", "Chef d'équipe"],
        # Poste a APOSTROPHE, joue expres : Jinja l'echappe en « d&#39; », et
        # les assertions d'ecran comparaient au texte brut -> faux rouge. Les
        # intitules francais courants en sont pleins (chef d'equipe, agent
        # d'entretien), aucun client factice n'en jouait un.
        "poste": "Chef d'équipe",
        "societes": [societe(
            "Toitures du Nord", "444 444 444", "compta@toitures.example",
            "TOITURES DU NORD SAS", "7 rue des Ardoises, 59000 Lille", [
                etab("Chantier Nord", "444 444 444 00015", "7 rue des Ardoises, 59000 Lille", "genere"),
                etab("Chantier Sud", "444 444 444 00023", "3 quai de la Deûle, 59000 Lille", "depose"),
            ])],
        "attendus": ["TOITURES DU NORD SAS", "444 444 444 00015", "Chef d'équipe"],
        "etab_genere": "Toitures du Nord / Chantier Nord",
        "etab_depose": "Toitures du Nord / Chantier Sud",
    },
]


def modele_docx(chemin, lignes):
    """Un .docx de test genere a la volee. Les fixtures ne peuvent pas venir de
    config/ : c'est l'instance locale, hors du depot depuis que le produit et
    l'instance sont separes -- un clone frais n'en a pas."""
    doc = Document()
    for l in lignes:
        doc.add_paragraph(l)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(chemin))


def poser_code(copie):
    """Une copie du PRODUIT : le code + le squelette config.exemple/, jamais
    l'instance d'un autre client."""
    (copie / "templates").mkdir(parents=True)
    for f in CODE:
        shutil.copy2(RACINE / f, copie / f)
    for t in (RACINE / "templates").glob("*.html"):
        shutil.copy2(t, copie / "templates" / t.name)
    shutil.copytree(RACINE / "config.exemple", copie / "config.exemple")


def installer_copie(base, client):
    copie = base / client["slug"]
    poser_code(copie)

    out = subprocess.run([sys.executable, "installer.py", client["nom"], "."],
                         cwd=copie, capture_output=True, text=True, encoding="utf-8")
    assert out.returncode == 0, out.stderr
    mdp = re.search(r"rh / (\S+)", out.stdout)
    assert mdp, out.stdout
    client["password"] = mdp.group(1)

    # config propre au client
    (copie / "config" / "societes.json").write_text(
        json.dumps(client["societes"], ensure_ascii=False, indent=2), encoding="utf-8")
    if formulaire := client.get("formulaire"):
        (copie / "config" / "formulaire.json").write_text(
            json.dumps(formulaire, ensure_ascii=False, indent=2), encoding="utf-8")
    inst = json.loads((copie / "config" / "instance.json").read_text(encoding="utf-8"))
    inst["templates"] = {p: "contrat.docx" for p in client.get(
        "postes", ("Équipier polyvalent", "Assistant manager", "Manager"))}
    inst["fiche_salarie"] = "fiche_salarie.docx"
    (copie / "config" / "instance.json").write_text(
        json.dumps(inst, ensure_ascii=False, indent=2), encoding="utf-8")
    modele_docx(copie / "config" / "contrats" / "contrat.docx", [
        "CONTRAT DE TRAVAIL",
        "Entre {{RaisonSociale}}, {{FormeCapital}}, {{RCS}},",
        "dont le siege social est {{SiegeSocial}},",
        "etablissement {{Etablissement}} — SIRET {{Siret}},",
        "et {{NomPrenom}}, engage(e) au poste de {{Poste}}.",
        "Convention collective : {{ConventionCollective}}.",
    ])
    modele_docx(copie / "config" / "contrats" / "fiche_salarie.docx", [
        "FICHE SALARIE — {{NomPrenom}}",
        "Pieces fournies : {{PiecesFournies}}",
        "Pieces manquantes : {{PiecesManquantes}}",
    ])
    return copie


def refus_avant_config(base, client):
    """Juste après installer.py, sans societes.json rempli : refus de servir."""
    copie = base / (client["slug"] + "-nu")
    poser_code(copie)
    subprocess.run([sys.executable, "installer.py", client["nom"], "."],
                   cwd=copie, capture_output=True, text=True, encoding="utf-8", check=True)
    assert subprocess.run([sys.executable, "installer.py", client["nom"], "."],
                          cwd=copie, capture_output=True, text=True,
                          encoding="utf-8").returncode != 0, \
        "réinstaller par-dessus une instance existante doit être refusé"
    r = subprocess.run(
        [sys.executable, "-c", "import config,sys; sys.exit(0 if config.verifier() else 1)"],
        cwd=copie, capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, "une copie fraîchement installée devrait refuser de servir"
    assert not list((copie / "config").rglob("*.docx")), "des .docx ont survécu à l'installation"
    for j in (copie / "config").rglob("*.json"):
        assert "wingstop" not in j.read_text(encoding="utf-8").lower(), f"« wingstop » dans {j.name}"


def main():
    base = Path(tempfile.mkdtemp(prefix="contratgen-clients-"))
    driver = str(Path(__file__).resolve().parent / "_client_parcours.py")
    vus = []
    try:
        for client in CLIENTS:
            refus_avant_config(base, client)
            copie = installer_copie(base, client)

            spec = base / f"{client['slug']}-spec.json"
            spec.write_text(json.dumps({
                "password": client["password"],
                "poste": client.get("poste", "Manager"),
                "attendus": client["attendus"],
                "etab_genere": client["etab_genere"],
                "etab_depose": client["etab_depose"],
            }, ensure_ascii=False), encoding="utf-8")

            env = {**os.environ, "DONNEES": str(copie / "data"),
                   "PYTHONIOENCODING": "utf-8"}
            r = subprocess.run([sys.executable, driver, str(copie), str(spec)],
                               capture_output=True, text=True, encoding="utf-8", env=env)
            assert r.returncode == 0, f"[{client['nom']}]\n{r.stdout}\n{r.stderr}"
            ok = [l for l in r.stdout.splitlines() if l.startswith("PARCOURS OK")]
            assert ok, r.stdout
            _, nom, secret8, donnees = ok[0].split("\t")
            print(f"  ✓ {nom:22} secret {secret8}…  {donnees}")
            vus.append((nom, secret8, donnees))

        noms = {n for n, _, _ in vus}
        secrets = {s for _, s, _ in vus}
        dossiers = {d for _, _, d in vus}
        assert len(noms) == len(secrets) == len(dossiers) == len(CLIENTS), \
            f"deux clients partagent un secret ou un dossier de données : {vus}"
        print(f"\n{len(CLIENTS)} clients factices — copies étanches, "
              f"secrets et données distincts.\n{base}")
    finally:
        shutil.rmtree(base, ignore_errors=True)


if __name__ == "__main__":
    main()