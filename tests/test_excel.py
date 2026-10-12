"""« Compléter les dossiers » : les salariés importés du SharePoint complétés
depuis les exports Excel de la RH (excel.py) ou à la main dans un tableau.

    python tests/test_excel.py

L'Excel ne remplace jamais un champ rempli ; une ligne non rapprochée se
rattache à la main ; la fiche salarié se génère une fois les champs complétés.
Classeurs construits dans le test, données dans un dossier temporaire.
"""
import io
import json
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ["DONNEES"] = tempfile.mkdtemp(prefix="contratgen-excel-")

import config   # noqa: E402

SP = {"site": "https://acme.sharepoint.com/sites/rh", "racine": "SALARIES",
      "postes": {"Second managers": "Assistant Manager"}}
config.sharepoint = lambda: SP
_texte = lambda i, ph, **kw: {"id": i, "libelle": i, "type": "texte", "placeholder": ph, **kw}  # noqa: E731
_date = lambda i, ph, **kw: {"id": i, "libelle": i, "type": "date", "placeholder": ph, **kw}    # noqa: E731
_choix = lambda i, ph, o, **kw: {"id": i, "libelle": i, "type": "choix", "placeholder": ph,     # noqa: E731
                                 "options": o, **kw}
AUTRES = ["nationalite", "Autres"]
config._CACHE["formulaire.json"] = {"version": 1, "roles": {"nom": "nom_prenom", "email": "email"},
    "champs": [
        {"id": "etablissement", "libelle": "Établissement", "type": "etablissement", "requis": True},
        _choix("poste", "Poste", ["Assistant Manager", "Equipier Polyvalent", "Manager"], requis=True),
        _choix("civilite", "Civilite", ["Monsieur", "Madame"], requis=True),
        _texte("nom_prenom", "NomPrenom", requis=True),
        _date("date_naissance", "DateNaissance", requis=True),
        _texte("telephone", "Telephone", requis=True),
        {"id": "email", "libelle": "email", "type": "email", "placeholder": "Email"},
        _texte("adresse", "Domicile", requis=True),
        _date("date_embauche", "DateEmbauche"),
        _choix("type_contrat", "TypeContrat", ["CDI", "CDD"], requis=True),
        _choix("nationalite", None, ["Français", "Autres"]),
        _texte("nationalite_etrangere", "NationaliteEtrangere", requis_si=AUTRES),
        _choix("type_autorisation", "TypeAutorisation",
               ["Carte de résident", "Carte de séjour", "VLS-TS", "Récépissé de renouvellement"],
               requis_si=AUTRES),
        _date("date_fin_validite", "DateFinValidite", requis_si=AUTRES),
        _date("date_debut", "DateDebut", requis=True),
        _choix("temps_partiel", None, ["OUI", "NON"]),
        _choix("temps_travail", "TempsTravail", ["18H", "24H", "35H"]),
        _texte("num_secu", "NumSS", requis=True),
        {"id": "cni", "libelle": "CNI", "type": "piece_jointe", "role": "cni", "requis": True}]}
config._CACHE["instance.json"] = {k: v for k, v in config.instance().items() if k != "fiche_salarie"}

import app      # noqa: E402
import excel    # noqa: E402
import sharepoint  # noqa: E402
import store    # noqa: E402

ETAB, AUTRE = "ACME Restauration / Lille Grand Place", "ACME Restauration / Paris Opéra"


def xlsx(lignes, partagees=True):
    """Classeur minimal : la 1re feuille est sheet2.xml (rId7), une feuille
    leurre sheet1.xml vient avant dans le zip. Texte enrichi : chaque chaîne
    partagée coupée en deux runs."""
    ss, rows = [], []
    for r, ligne in enumerate(lignes, 1):
        cells = []
        for c, v in enumerate(ligne):
            ref = ("" if c < 26 else chr(64 + c // 26)) + chr(65 + c % 26) + str(r)
            if v is None:
                continue
            if callable(v):                     # cellule brute : formule, erreur...
                cells.append(v(ref))
                continue
            if isinstance(v, (int, float)):
                cells.append(f'<c r="{ref}"><v>{v}</v></c>')
            elif partagees:
                ss.append(f"<si><r><t>{escape(v[:2])}</t></r><r><t>{escape(v[2:])}</t></r></si>")
                cells.append(f'<c r="{ref}" t="s"><v>{len(ss) - 1}</v></c>')
            else:
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(v)}</t></is></c>')
        rows.append(f'<row r="{r}">{"".join(cells)}</row>')
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rns = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    feuille = lambda corps: f'<worksheet {ns}><sheetData>{corps}</sheetData></worksheet>'  # noqa: E731
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/worksheets/sheet1.xml", feuille(
            '<row r="1"><c r="A1" t="inlineStr"><is><t>Nom</t></is></c>'
            '<c r="B1" t="inlineStr"><is><t>Prénom</t></is></c></row>'
            '<row r="2"><c r="A2" t="inlineStr"><is><t>LEURRE</t></is></c>'
            '<c r="B2" t="inlineStr"><is><t>Max</t></is></c></row>'))
        z.writestr("xl/worksheets/sheet2.xml", feuille("".join(rows)))
        z.writestr("xl/workbook.xml", f'<workbook {ns} {rns}><sheets>'
                   '<sheet name="Liste" sheetId="1" r:id="rId7"/>'
                   '<sheet name="Autre" sheetId="2" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels",
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Target="worksheets/sheet1.xml" Type="x"/>'
                   '<Relationship Id="rId7" Target="/xl/worksheets/sheet2.xml" Type="x"/>'
                   '</Relationships>')
        if partagees:
            z.writestr("xl/sharedStrings.xml", f'<sst {ns}>{"".join(ss)}</sst>')
    return buf.getvalue()


# 1. lecture
EXPORT_RH = ["Nom", "Prénom", "Sexe", "Date de naissance", "N° Téléphone", "Email", "Adresse 1",
          "Complément", "Code Postal", "Ville", "Nationalité", "Fin de carte de séjour",
          "Nom de permis de travail / carte de séjour", "Date d'entrée", "Poste", "Type de contrat",
          "Contrat hebdo", "N° de sécurité sociale"]
classeur = xlsx([
    ["Liste des employés"],
    EXPORT_RH,
    ["DUPONT", "Jean", "Masculin", 32874, "06 11 22 33 44", "jean@x.fr", "1 rue A", "", "59000",
     "Lille", "Français-e", "N/A", "", 45678.0, "Equipier", "Durée indéterminée", 35.0,
     "1 85 05 75 123 456 78"],
    ["MARTIN", "Paul", "Masculin", 33000, "", "", "", "", "", "", "Indien-ne", "", "VLSTS",
     "29/05/25", "Second Manager", "Apprenti", 151.67],
    ["SANSPRENOM", "", "Feminin"],
])
lues = excel.lignes(classeur)
assert [(l["nom"], l["prenom"]) for l in lues] == [("DUPONT", "Jean"), ("MARTIN", "Paul")], lues
assert lues[0]["datedentree"] == "45678.0" and lues[0]["ntelephone"] == "06 11 22 33 44"
assert excel.lignes(xlsx([EXPORT_RH, ["A", "B"]], partagees=False))[0]["prenom"] == "B", "inlineStr"
planning = excel.lignes(xlsx([
    ["Nom", "Prénom", "Volume Horaire Hebdomadaire (h)*", "Type de Contrat", "Intitulé Poste"],
    ["FORFAITS JOUR"], ["LEROY", "Ana", 18.0, "CDI", "Equipier polyvalent"],
    ["Nom", "Prénom", "Type de Contrat", "Intitulé Poste"], ["EXTRAS"],
    ["BLANC", "Eve", "CDD", "Manager"]], partagees=False))
assert [(l["nom"], l["typedecontrat"], l["intituleposte"]) for l in planning] == [
    ("LEROY", "CDI", "Equipier polyvalent"), ("BLANC", "CDD", "Manager")], planning
for mauvais in (b"pas un zip", xlsx([EXPORT_RH])[:200]):
    try:
        excel.lignes(mauvais)
        raise AssertionError("fichier illisible accepté")
    except ValueError:
        pass
print("ok  lecture : 1re feuille par les rels, texte enrichi, inlineStr, sections, titres ignorés")

# 2. conversions
assert excel._date("45678.0") == "2025-01-21" and excel._date("29/05/25") == "2025-05-29"
assert excel._date("01/02/1990") == "1990-02-01" and excel._date("bientôt") is None
v, r = excel.convertir(lues[0])
assert v == {"civilite": "Monsieur", "date_naissance": "1990-01-01", "telephone": "06 11 22 33 44",
             "email": "jean@x.fr", "adresse": "1 rue A, 59000 Lille", "nationalite": "Français",
             "date_debut": "2025-01-21", "date_embauche": "2025-01-21", "type_contrat": "CDI",
             "poste": "Equipier Polyvalent", "temps_travail": "35H", "temps_partiel": "NON",
             "num_secu": "185057512345678"}, v
assert r == {} and "date_fin_validite" not in v, "N/A converti"
v, r = excel.convertir(lues[1])
assert (v["nationalite"], v["nationalite_etrangere"], v["type_autorisation"]) == ("Autres", "Indien", "VLS-TS")
assert v["poste"] == "Assistant Manager", "alias sharepoint.postes"
assert r == {"type_contrat": "Apprenti", "temps_travail": "151.67"} and "temps_partiel" not in v, r
SP["postes"] = {}
assert excel.convertir(lues[1])[1]["poste"] == "Second Manager", "« Second Manager » devenu « Manager »"
SP["postes"] = {"Second managers": "Assistant Manager"}
cas = lambda **l: excel.convertir({"nom": "X", "prenom": "Y", **l})  # noqa: E731
v = cas(sexe="Feminin", nationalite="Indien-ne")[0]
assert (v["civilite"], v["nationalite"], v["nationalite_etrangere"]) == ("Madame", "Autres", "Indienne")
assert cas(sexe="Feminin", nationalite="Bangladais-e, Bangladeshi-e")[0]["nationalite_etrangere"] == "Bangladaise"
assert cas(nationalite="Français-e")[0]["nationalite"] == "Français"
OPT = lambda brut: cas(nomdepermisdetravailcartedesejour=brut)  # noqa: E731
assert OPT("Carte de séjour pluriannuelle")[0]["type_autorisation"] == "Carte de séjour"
assert OPT("Attestation de prolongation")[1]["type_autorisation"] == "Attestation de prolongation"
assert OPT("Carte inconnue")[1], "1er mot commun à deux options : ambigu"
assert cas(poste="Equipier")[0]["poste"] == "Equipier Polyvalent"
assert cas(typedecontrat="Durée déterminée")[0]["type_contrat"] == "CDD"
v = cas(contrathebdo="18.0")[0]
assert (v["temps_travail"], v["temps_partiel"]) == ("18H", "OUI")
print("ok  conversions : dates, sexe, nationalité accordée, permis, postes, contrat, heures")

# 3. fusion de plusieurs fichiers
f1 = xlsx([EXPORT_RH[:3] + ["Poste"], ["DUPONT", "Jean", "Masculin", "Manager"]])
f2 = xlsx([["Prénom", "Nom", "Date d'entrée", "Intitulé Poste"], ["Jean", "Dupont", "01/03/24", "Equipier"]])
p = excel.par_personne([("a.xlsx", f1), ("b.xlsx", f2)])[sharepoint.cle_nom("Jean DUPONT")]
assert p["valeurs"]["civilite"] == "Monsieur" and p["valeurs"]["date_debut"] == "2024-03-01"
assert p["conflits"] == {"poste": ["Equipier Polyvalent", "Manager"]} and "poste" not in p["valeurs"]
assert p["sources"] == ["a.xlsx", "b.xlsx"]
f3 = xlsx([["Nom", "Prénom", "Poste"], ["DUPONT", "Jean", "Directeur"]])
p = excel.par_personne([("a.xlsx", f1), ("c.xlsx", f3)])[sharepoint.cle_nom("Jean DUPONT")]
assert p["conflits"]["poste"] == ["Directeur", "Manager"], "rejet masqué par une valeur d'un autre fichier"
assert "poste" not in p["valeurs"], "valeur écrite malgré le désaccord"
try:
    excel.par_personne([("vieux.xls", b"\xd0\xcf")])
    raise AssertionError(".xls accepté")
except ValueError as e:
    assert "vieux.xls" in str(e), e
print("ok  fusion : fichiers complémentaires réunis, désaccord = conflit, jamais rempli")

# --- dossiers importés ----------------------------------------------------------


def importe(nom, etab=ETAB, **champs):
    return store.importer({"etablissement": etab, "nom_prenom": nom, "poste": "Manager", **champs},
                          "sp:" + nom + etab, {"par": "rh", "origine": "sharepoint"}, [])


dupont = importe("DUPONT Jean", telephone="0700000000")
martin1, martin2 = importe("MARTIN Paul"), importe("MARTIN Paul (2)")
store.completer_champs(martin2, {"nom_prenom": "MARTIN Paul"})           # vrais homonymes
ailleurs = importe("DUPONT Jean", etab=AUTRE)
leroy = importe("LEROY Ana")
soumis = store.creer_soumission({"etablissement": ETAB, "nom_prenom": "DUPONT Jean"}, {})

app.app.config["SESSION_COOKIE_SECURE"] = False
c = app.app.test_client()
c.post("/login", data={"identifiant": "rh", "mot_de_passe": "fixture"})
import html as _html  # noqa: E402

page = lambda r: _html.unescape(r.get_data(as_text=True))  # noqa: E731

# 8. GET
assert c.get("/etablissements/completer?cle=Inconnu").status_code == 404
s = page(c.get(f"/etablissements/completer?cle={ETAB}"))
assert f"/dossier/{dupont}" in s and f"/dossier/{leroy}" in s
assert f"/dossier/{ailleurs}" not in s and f"/dossier/{soumis}" not in s
assert "Compléter les dossiers" in page(c.get("/etablissements"))
print("ok  tableau : importés en poste de l'établissement seulement")

# 4. upload : propose, n'écrit rien avant « Enregistrer »
envoi = lambda *fs: c.post("/etablissements/completer/excel", content_type="multipart/form-data",  # noqa: E731
                           data={"cle": ETAB, "fichiers": [(io.BytesIO(b), n) for n, b in fs]})


def post(**d):
    """POST du tableau, avec la version affichée de chaque ligne (verrou optimiste)."""
    v = {f"v:{i['id']}": app._version(i["champs"]) for i in app._importes(ETAB)}
    return c.post("/etablissements/completer", data={"cle": ETAB, **v, **d})


def formulaire(s):
    """Ce que le navigateur posterait : le tableau tel qu'il est affiché."""
    s = s.split('id="tableau-completer"')[1]
    d = dict(re.findall(r'<input type="(?:text|hidden)" name="([^"]+)" value="([^"]*)"', s))
    for m in re.finditer(r'<select name="([^"]+)"[^>]*>(.*?)</select>', s, re.S):
        sel = re.search(r'<option(?: value="([^"]*)")? selected>([^<]*)', m.group(2))
        d[m.group(1)] = (sel.group(1) if sel.group(1) is not None else sel.group(2)) if sel else ""
    return d


avant = dict(store.lire(dupont)["champs"])
r = envoi(("export.xlsx", classeur))
s = page(r)
assert r.status_code == 200 and "1 dossier(s) rapproché(s)" in s and "Rien n'est enregistré" in s, s[:3000]
assert store.lire(dupont)["champs"] == avant, "l'Excel a écrit sans « Enregistrer »"
assert "proposé par l'Excel" in s and 'value="01/01/1990"' in s, "valeur proposée absente"
assert re.search(r'<td class="propose', s) and re.search(r'<td class="ecart', s), "cellules non surlignées"
assert "Excel : 06 11 22 33 44" in s and "Excel : Equipier Polyvalent" in s, "écart non signalé"
assert "MARTIN Paul" in s and "non rattachées" in s, "homonymes côté dossiers"
assert "2 dossiers en poste à ce nom" in s, "motif de l'homonyme absent"
assert post(**formulaire(s)).status_code == 302                     # la RH valide
d = store.lire(dupont)["champs"]
assert d["telephone"] == "0700000000" and d["poste"] == "Manager", "champ rempli remplacé"
assert d["civilite"] == "Monsieur" and d["date_naissance"] == "1990-01-01" and d["num_secu"]
assert not store.lire(martin1)["champs"].get("civilite") and not store.lire(martin2)["champs"].get("civilite")
assert not store.lire(ailleurs)["champs"].get("civilite"), "autre établissement touché"
assert "civilite" not in store.lire(soumis)["champs"], "dossier non importé touché"
entree = store.lire(dupont)["journal"][-1]
assert entree["type"] == "correction" and "civilite" in entree["champs"]
assert "Monsieur" not in json.dumps(entree, ensure_ascii=False) and "1990" not in json.dumps(entree)
n = len(store.lire(dupont)["journal"])
s = page(envoi(("export.xlsx", classeur)))
assert "0 champ(s) proposé(s)" in s
assert post(**formulaire(s)).status_code == 302 and len(store.lire(dupont)["journal"]) == n, \
    "renvoi : entrée de journal vide"
r = envoi(("faux.xlsx", b"rien"))
assert r.status_code == 302 and "faux.xlsx" in page(c.get(r.location))
print("ok  upload : propose sans écrire, surligne, la RH enregistre ; remplis intacts, journal sans valeurs")

# 5. tableau
r = post(**{f"{leroy}:date_naissance": "01/02/1990", f"{leroy}:telephone": "0611",
            f"{ailleurs}:civilite": "Madame"})
assert r.status_code == 302
assert store.lire(leroy)["champs"]["date_naissance"] == "1990-02-01"
assert not store.lire(ailleurs)["champs"].get("civilite"), "uid étranger écrit"
assert store.lire(leroy)["journal"][-1]["type"] == "correction"
assert post(**{f"{leroy}:telephone": ""}).status_code == 302
assert store.lire(leroy)["champs"].get("telephone") == "", "une cellule vidée doit effacer"
assert "telephone" in store.lire(leroy)["journal"][-1]["champs"]
config.champ("telephone")["libelle"] = "Téléphone portable"      # libellé ≠ id, le temps de lire
s = page(c.get(f"/dossier/{leroy}"))
config.champ("telephone")["libelle"] = "telephone"
assert re.search(r"<b>correction</b>\s*par rh\s*· Téléphone portable\s*</li>", s), \
    "journal de la fiche : libellés des champs corrigés non affichés"
le = store.lire(leroy)["journal"][-1]["le"]
assert f"<time>{app._horodatage(le)}</time>" in s, "journal : heure UTC brute au lieu de l'heure locale"
assert app._horodatage("pas une date") == "pas une date"
r = post(**{f"{leroy}:date_naissance": "31/02/1990", f"{leroy}:telephone": "0611"})
s = page(r)
assert r.status_code == 422 and "31/02/1990" in s
assert re.search(r'<td class="erreur[^"]*">\s*<input[^>]*value="31/02/1990"[^>]*aria-invalid', s), \
    "cellule fautive non surlignée"
r = post(**{f"{leroy}:civilite": "Mademoiselle", f"{leroy}:telephone": "0611"})
assert r.status_code == 422 and not store.lire(leroy)["champs"].get("telephone"), "écrit malgré l'erreur"
r = post(**{f"{leroy}:telephone": "0622", f"v:{leroy}": "1"})          # page affichée avant une autre saisie
assert r.status_code == 422 and "modifié ailleurs" in page(r)
assert not store.lire(leroy)["champs"].get("telephone"), "ligne périmée écrite"
r = c.post("/etablissements/completer", data={"cle": ETAB, f"{leroy}:telephone": "0622"})
assert r.status_code == 422, "sans version : écriture à l'aveugle"
# deux onglets : la page réaffichée après un refus ne doit pas défaire l'autre onglet
vue = formulaire(page(c.get(f"/etablissements/completer?cle={ETAB}")))
assert post(**{f"{leroy}:telephone": "0633"}).status_code == 302            # 1er onglet
r = c.post("/etablissements/completer", data={**vue, f"{leroy}:date_naissance": "02/02/1990"})
assert r.status_code == 422 and "modifié ailleurs" in page(r)
assert c.post("/etablissements/completer", data=formulaire(page(r))).status_code == 302
assert store.lire(leroy)["champs"]["telephone"] == "0633", "ré-enregistrer a défait le 1er onglet"
vue = formulaire(page(c.get(f"/etablissements/completer?cle={ETAB}")))
assert post(**{f"{leroy}:telephone": "0644"}).status_code == 302            # 1er onglet
r = c.post("/etablissements/completer", data={**vue, f"{dupont}:date_naissance": "31/02/1990"})
assert r.status_code == 422                                                 # erreur sur une autre ligne
corrige = {**formulaire(page(r)), f"{dupont}:date_naissance": vue[f"{dupont}:date_naissance"]}
assert c.post("/etablissements/completer", data=corrige).status_code == 302
assert store.lire(leroy)["champs"]["telephone"] == "0644", "ligne non touchée réécrite, périmée"
store.archiver(martin2, "rh", "2026-09-30")
try:
    store.completer(martin2, {"civilite": "Monsieur"}, "rh", "correction")
    raise AssertionError("dossier Parti complété")
except ValueError:
    pass
assert store.date_sortie(store.lire(martin2)) == "2026-09-30"
print("ok  tableau : ISO, vider efface, erreur surlignée = 422 sans écriture, ligne modifiée ailleurs refusée")

# 6. rattachement manuel
r = post(**{"x0": dupont, "x0:_nom": "DUPONT J.", "x0:telephone": "0600", "x0:type_contrat": "CDD",
            "x1": "", "x1:civilite": "Madame"})
assert r.status_code == 302
d = store.lire(dupont)["champs"]
assert d["telephone"] == "0700000000" and d["type_contrat"] == "CDI", "rattachement a remplacé"
store.completer(martin1, {}, "rh", "x")                                   # sans effet
r = post(**{"x0": martin1, "x0:_nom": "MARTIN Paul", "x0:civilite": "Monsieur",
            "x0:date_naissance": "1990-05-05"})
assert store.lire(martin1)["champs"]["civilite"] == "Monsieur"
assert store.lire(martin1)["champs"]["date_naissance"] == "1990-05-05"
print("ok  rattachement : seuls les champs vides du dossier choisi")

# 7. fiche salarié
r = c.post(f"/dossier/{leroy}/fiche")
motif = store.lire(leroy)["journal"][-1]["motif"]
assert "Fiche non générée" in page(c.get(r.location)) and "civilite" in motif, motif
assert "Civilite" not in motif, "jeton du modèle au lieu du libellé du champ"
post(**{f"{leroy}:{k}": v for k, v in {
    "civilite": "Madame", "date_debut": "01/03/2024", "adresse": "2 rue B", "num_secu": "2900159",
    "telephone": "0612", "type_contrat": "CDI"}.items()})
fiche = store.nom_piece(store.lire(leroy), "fiche-salarie", ".docx")
r = c.post(f"/dossier/{leroy}/fiche")
assert "Fiche salarié générée" in page(c.get(r.location)), store.lire(leroy)["journal"][-1]
assert fiche in store.fichiers(store.lire(leroy), "pieces")
from docx import Document  # noqa: E402

texte = "\n".join(p.text for p in Document(io.BytesIO(store.ouvrir(leroy, "pieces", fiche))).paragraphs)
texte += "\n".join(cl.text for t in Document(io.BytesIO(store.ouvrir(leroy, "pieces", fiche))).tables
                   for row in t.rows for cl in row.cells)
assert "LEROY" in texte and "février 1990" in texte and "non vérifiées (dossier importé)" in texte, texte
c.post(f"/dossier/{leroy}/fiche")
assert store.fichiers(store.lire(leroy), "pieces").count(fiche) == 1
orig_nom = store.nom_piece({"champs": {"nom_prenom": "DURAND Luc"}}, "fiche-salarie", ".docx")
durand = store.importer({"etablissement": ETAB, "nom_prenom": "DURAND Luc", "civilite": "Monsieur",
                         "date_naissance": "1980-01-01", "poste": "Manager"}, "sp:DURAND",
                        {"par": "rh"}, [("pieces", orig_nom, lambda: b"original")])
r = c.post(f"/dossier/{durand}/fiche")
assert "n'est pas remplacé" in page(c.get(r.location))
assert store.ouvrir(durand, "pieces", orig_nom) == b"original"
assert c.post(f"/dossier/{soumis}/fiche").status_code == 400
assert "Générer la fiche salarié" in page(c.get(f"/dossier/{leroy}"))
print("ok  fiche : refusée avec le motif, générée puis régénérée, fichier d'origine protégé")

# 9. cas limites -----------------------------------------------------------------
brut = lambda xml: (lambda ref: xml.format(ref=ref))  # noqa: E731

# 9a. lecture : formule, erreur, colonnes au-delà de Z, fichiers malformés
large = ["Nom", "Prénom"] + [f"Vide{i}" for i in range(26)] + ["Sexe"]
l = excel.lignes(xlsx([
    ["Nom", "Prénom", "Date de naissance", "Poste", "Contrat hebdo"],
    ["FORM", "Ule", brut('<c r="{ref}" t="e"><v>#N/A</v></c>'),
     brut('<c r="{ref}" t="str"><f>A1</f><v>Manager</v></c>'), brut('<c r="{ref}"><f>1+1</f></c>')],
    large, ["LOIN", "Col"] + [""] * 26 + ["Feminin"]], partagees=False))
assert l[0]["datedenaissance"] == "#N/A" and l[0]["poste"] == "Manager" and l[0]["contrathebdo"] == ""
v, r = excel.convertir(l[0])
assert v == {"poste": "Manager"} and r == {"date_naissance": "#N/A"}, (v, r)
assert l[1]["sexe"] == "Feminin", "colonne AC"
for nom, malforme in [
        ("index de chaîne partagée", xlsx([EXPORT_RH, ["A", "B"]]).replace(b"<v>1</v>", b"<v>99</v>")),
        ("cellule sans r", xlsx([["Nom", "Prénom"], [brut('<c t="inlineStr"><is><t>X</t></is></c>'), "B"]],
                                partagees=False)),
        ("index non numérique", xlsx([EXPORT_RH, ["A", "B"]]).replace(b"<v>1</v>", b"<v>x</v>"))]:
    try:
        excel.lignes(malforme)
        raise AssertionError(f"{nom} : accepté")
    except ValueError as e:
        assert str(e) == "classeur .xlsx illisible", (nom, e)          # sans nom d'exception Python
vide = excel.lignes(xlsx([["Rien d'utile"]]))
assert vide == [] and excel.par_personne([("vide.xlsx", xlsx([EXPORT_RH]))]) == {}
print("ok  lecture : formule, #N/A, colonne AC, classeurs malformés refusés proprement, feuille vide")

# 9b. conversions aux bords
assert excel._date("2024-03-01") == "2024-03-01", "date déjà ISO"
assert excel._date("99999999") is None and excel._date("0") is None, "nombre aberrant pris pour une date"
assert excel._date("31/02/2024") is None and excel._date("01/13/2024") is None
v, r = cas(sexe="Inconnu", nationalite="Indien-ne")
assert r["civilite"] == "Inconnu" and v["nationalite_etrangere"] == "Indien-ne", "accord sans sexe connu"
assert cas(nationalite="Française")[0]["nationalite"] == "Français"
assert cas(nationalite="Belge", sexe="Feminin")[0]["nationalite_etrangere"] == "Belge"
assert cas(contrathebdo="35,5")[1] == {"temps_travail": "35,5"}
assert cas(contrathebdo="0.0")[1] == {"temps_travail": "0.0"}
assert cas(contrathebdo="24")[0]["temps_travail"] == "24H"
v = cas(adresse1="3 rue C", ville="Lyon")[0]
assert v["adresse"] == "3 rue C, Lyon", v
assert cas(email="", emailsecondaire="b@x.fr")[0]["email"] == "b@x.fr"
assert cas(datedentree="", datedanciennete="N/A", datedebutdecontrat="01/01/24")[0]["date_debut"] == "2024-01-01"
assert cas(poste="")[0] == {} and cas()[0] == {}, "cellule vide convertie"
print("ok  conversions : ISO, nombre aberrant, sexe inconnu, nationalités, heures, repli de colonnes")

# 9c. fusion : doublons identiques, ordre et accents des noms, nom à 3 mots
f4 = xlsx([["Nom", "Prénom", "Sexe"], ["Dupont", "Jean", "Masculin"], ["DUPONT", "JEAN", "Masculin"]])
pp = excel.par_personne([("a.xlsx", f4), ("a (1).xlsx", f4)])
assert len(pp) == 1 and next(iter(pp.values()))["conflits"] == {}, "doublons identiques = conflit"
assert sharepoint.cle_nom("Élodie", "MARTIN-DURAND") == sharepoint.cle_nom("martin durand elodie")
assert sharepoint.cle_nom("DURAND", "Jean Marc") != sharepoint.cle_nom("DURAND Jean"), "3 mots vs 2"
print("ok  fusion : doublons identiques sans conflit, accents/ordre/tirets, prénoms en plus non rapprochés")

# 9d. accès, garde-fous de route
anonyme = app.app.test_client()
for u in (f"/etablissements/completer?cle={ETAB}",):
    assert anonyme.get(u).status_code == 302
assert anonyme.post("/etablissements/completer", data={"cle": ETAB}).status_code == 302
assert anonyme.post(f"/dossier/{leroy}/fiche").status_code == 302
assert c.get("/etablissements/completer").status_code == 404, "sans clé"
config.sharepoint = lambda: {}
assert c.get(f"/etablissements/completer?cle={ETAB}").status_code == 404
assert "Compléter les dossiers" not in page(c.get("/etablissements"))
config.sharepoint = lambda: SP
r = c.post("/etablissements/completer/excel", data={"cle": ETAB}, content_type="multipart/form-data")
assert r.status_code == 302 and "Choisissez au moins un fichier" in page(c.get(r.location))
r = envoi(("cv.pdf", b"%PDF-1.4"))
assert r.status_code == 302 and "cv.pdf" in page(c.get(r.location))
s = page(c.get(f"/etablissements/completer?cle={AUTRE}"))
assert f"/dossier/{ailleurs}" in s and f"/dossier/{dupont}" not in s
print("ok  accès : connexion exigée, 404 sans clé ou sans bloc sharepoint, upload vide ou PDF refusé")

# 9e. upload : conflit et rejet affichés sous la cellule, jamais écrits ; ligne d'un Parti
nouveau = importe("NOUVEAU Zoe")
conflit = xlsx([["Nom", "Prénom", "Poste", "Type de contrat"], ["NOUVEAU", "Zoe", "Manager", "Apprenti"]])
conflit2 = xlsx([["Nom", "Prénom", "Poste"], ["NOUVEAU", "Zoe", "Equipier"]])
store.completer(nouveau, {}, "rh", "x")
ch = dict(store.lire(nouveau)["champs"])
s = page(envoi(("c1.xlsx", conflit), ("c2.xlsx", conflit2)))
assert "Excel : Equipier Polyvalent / Manager" in s, "conflit invisible"
assert "Excel : Apprenti (non reconnu)" in s, "rejet invisible"
assert store.lire(nouveau)["champs"] == ch, "conflit ou rejet écrit"
s = page(envoi(("parti.xlsx", xlsx([["Nom", "Prénom", "Sexe"], ["MARTIN", "Paul", "Masculin"]]))))
assert "1 dossier(s) rapproché(s)" in s, "homonyme parti : le dossier en poste est le seul candidat"
store.archiver(seul := importe("PARTI Seul"), "rh", "2026-09-30")
s = page(envoi(("partis.xlsx", xlsx([["Nom", "Prénom", "Sexe"], ["PARTI", "Seul", "Masculin"],
                                     ["INCONNU", "Max", "Masculin"]]))))
assert "1 ligne(s) Excel non rattachée(s)" in s and "1 ligne(s) de salarié parti ignorée(s)" in s, s[:2000]
assert "PARTI Seul" not in s.split("non rattachées")[1], "parti proposé au rattachement : erreur de dossier"
assert "INCONNU Max" in s and "aucun dossier en poste à ce nom" in s, "motif de l'inconnu absent"
r = post(**{"x0": leroy, "x0:_nom": "INCONNU Max", "x0:_motif": "aucun dossier en poste à ce nom",
            "x1": leroy, "x1:_nom": "X"})
assert r.status_code == 422 and "aucun dossier en poste à ce nom" in page(r), "motif perdu au réaffichage"
print("ok  upload : conflit et rejet signalés « Excel : … » sans écriture, homonyme parti écarté, "
      "parti seul ignoré, motif par ligne")

# 9f. tableau : champs hors colonnes, espaces, inchangé hors format, hors liste affiché, tableau entier
r = post(**{f"{leroy}:nom_prenom": "PIRATE", f"{leroy}:etablissement": AUTRE, f"{leroy}:cni": "x.pdf",
            f"{leroy}:telephone": "   "})
d = store.lire(leroy)["champs"]
assert d["nom_prenom"] == "LEROY Ana" and d["etablissement"] == ETAB and d["telephone"] == "", "espaces = vide"
with store._verrou(nouveau):                                 # valeurs anciennes hors format
    it = store.lire(nouveau)
    it["champs"].update(date_naissance="1er mai 1990", poste="Assistant manager")
    store._ecrire(store._rep(it), it)
s = page(c.get(f"/etablissements/completer?cle={ETAB}"))
assert "Assistant manager (hors liste)" in s, "choix hors liste affiché vide"
assert 'value="1er mai 1990"' in s
toutes = formulaire(s)
assert sum(1 for k in toutes if re.fullmatch(r"[0-9A-Z]{26}:[a-z_]+", k)) == len(app._importes(ETAB)) * len(app._colonnes())
r = post(**toutes)                                           # le formulaire tel quel, sans retouche
assert r.status_code == 302, "tableau renvoyé tel quel refusé (valeur ancienne hors format)"
assert store.lire(nouveau)["champs"]["date_naissance"] == "1er mai 1990"
toutes[f"{nouveau}:date_naissance"] = "02/05/1990"
assert post(**toutes).status_code == 302 and store.lire(nouveau)["champs"]["date_naissance"] == "1990-05-02"
gros = {f"{dupont}:telephone": "0700000000", **{f"x{n}": "" for n in range(1500)}}
assert post(**gros).status_code == 302, "plus de 1000 champs urlencoded"
print("ok  tableau : nom/établissement/pièce ignorés, espaces = vide, vieille valeur ne bloque pas, >1000 champs")

# 9g. rattachement : dossier d'un autre établissement, index non numérique, valeur forgée
r = post(**{"x0": ailleurs, "x0:civilite": "Madame", "xa": dupont, "xa:civilite": "Madame"})
assert r.status_code == 302 and not store.lire(ailleurs)["champs"].get("civilite")
r = post(**{"x0": nouveau, "x0:_nom": "NOUVEAU Z.", "x0:civilite": "Reine"})
assert r.status_code == 422 and "NOUVEAU Z." in page(r) and not store.lire(nouveau)["champs"].get("civilite")
print("ok  rattachement : autre établissement ignoré, valeur forgée = 422 réaffiché")

# 9h. fiche : parti, purgé, inconnu ; établissement retiré de la config
assert c.post(f"/dossier/{martin2}/fiche").status_code == 400, "fiche sur un Parti"
assert c.post("/dossier/INCONNU/fiche").status_code == 404
with store._verrou(martin1):
    it = store.lire(martin1)
    it["purge"] = {"le": store.maintenant(), "jours": 0}
    store._ecrire(store._rep(it), it)
assert c.post(f"/dossier/{martin1}/fiche").status_code == 400, "fiche sur un purgé"
assert f"/dossier/{martin1}" not in page(c.get(f"/etablissements/completer?cle={ETAB}")), "purgé listé"
vrai_mentions = config.mentions
config.mentions = lambda cle: (_ for _ in ()).throw(KeyError(cle))
r = c.post(f"/dossier/{leroy}/fiche")
config.mentions = vrai_mentions
assert "absent de config/societes.json" in page(c.get(r.location))
print("ok  fiche : 400 sur Parti ou purgé, 404 inconnu, établissement retiré = motif lisible")

# 10. tests poussés et cas d'erreur ------------------------------------------------
import time  # noqa: E402

GARCIA = importe("GARCIA Lou")
DARC = importe("D'ARC Jeanne")
brut_page = lambda r: r.get_data(as_text=True)  # noqa: E731

# 10a. version d'une ligne : une note (fiche) ne la périme pas, une saisie si
v_avant = app._version(store.lire(GARCIA)["champs"])
store.noter(GARCIA, type="fiche_echouee", par="systeme", motif="test")
assert post(**{f"{GARCIA}:telephone": "0601", f"v:{GARCIA}": v_avant}).status_code == 302, \
    "une simple note du journal a périmé la ligne"
r = post(**{f"{GARCIA}:telephone": "0602", f"v:{GARCIA}": v_avant})
assert r.status_code == 422 and store.lire(GARCIA)["champs"]["telephone"] == "0601"
assert post(**{f"{GARCIA}:telephone": "0602", f"v:{GARCIA}": "n'importe quoi"}).status_code == 422
print("ok  version : note sans effet, saisie concurrente refusée, version forgée refusée")

# 10b. injection HTML depuis l'Excel ou le tableau : toujours échappée
piege = xlsx([["Nom", "Prénom", "Adresse 1"], ["GARCIA", "Lou", '<script>alert(1)</script>'],
              ["<b>GRAS</b>", "<i>x</i>", "1 rue"]])
r = envoi(("piege.xlsx", piege))
t = brut_page(r)
assert "<script>alert(1)</script>" not in t and "&lt;script&gt;" in t, "script de l'Excel non échappé"
assert "<b>GRAS</b>" not in t and "&lt;b&gt;GRAS" in t, "nom de l'Excel non échappé"
post(**{f"{GARCIA}:adresse": '"><script>x</script>'})
t = brut_page(c.get(f"/etablissements/completer?cle={ETAB}"))
assert '"><script>x</script>' not in t, "valeur stockée réinjectée telle quelle"
post(**{f"{GARCIA}:adresse": "1 rue Sûre"})
print("ok  injection : valeurs de l'Excel et du tableau échappées à l'affichage")

# 10c. bombe zip et gros classeur
buf = io.BytesIO()
with zipfile.ZipFile(io.BytesIO(xlsx([EXPORT_RH]))) as src, \
        zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
    for n in src.namelist():
        # le vrai contenu + 25 Mo d'espaces : un XML valide, que seul le plafond refuse
        z.writestr(n, src.read(n) if n != "xl/sharedStrings.xml"
                   else src.read(n).replace(b"</sst>", b" " * (25 * 2**20) + b"</sst>"))
t0 = time.time()
try:
    excel.lignes(buf.getvalue())
    raise AssertionError("membre de 25 Mo décompressé")
except ValueError:
    pass
assert time.time() - t0 < 2 and len(buf.getvalue()) < 2**20, "bombe zip lue avant d'être refusée"
gros_classeur = xlsx([EXPORT_RH] + [[f"NOM{i}", f"P{i}", "Masculin", 33000 + i] for i in range(5000)],
                     partagees=False)
t0 = time.time()
assert len(excel.par_personne([("gros.xlsx", gros_classeur)])) == 5000
assert time.time() - t0 < 10, f"5000 lignes en {time.time() - t0:.1f} s"
print(f"ok  bombe zip refusée sans décompresser, 5000 lignes lues en {time.time() - t0:.1f} s")

# 10d. en-têtes et noms : majuscules, colonne en double, espace insécable, apostrophe
l = excel.lignes(xlsx([["NOM", "PRENOM", "SEXE", "Nom"], ["DUPONT\xa0", "\xa0Jean", "Masculin", "AUTRE"]]))
assert l[0]["nom"] == "DUPONT" and l[0]["prenom"] == "Jean" and l[0]["sexe"] == "Masculin", l
s = page(envoi(("darc.xlsx", xlsx([["Nom", "Prénom", "Sexe"], ["D'ARC", "Jeanne", "Feminin"]]))))
assert "1 dossier(s) rapproché(s)" in s, "nom avec apostrophe non rapproché"
print("ok  en-têtes en capitales, 1re colonne en double gardée, espaces insécables, apostrophe")

# 10e. NIR
for brut_nir, attendu in [("1 85 05 75 123 456 78", "185057512345678"), ("185057512345678", "185057512345678"),
                          ("2 90 05 2A 123 456", "290052A123456"), ("7 99 99 99 999 999 99", "799999999999999"),
                          ("1.85057512345678E+14", None), ("123", None), ("ABCDEFGHIJKLM", None)]:
    v, r = cas(ndesecuritesociale=brut_nir)
    assert v.get("num_secu") == attendu and (attendu or r.get("num_secu") == brut_nir), (brut_nir, v, r)
print("ok  NIR : espaces retirés, Corse, provisoire ; notation scientifique et formes fausses rejetées")

# 10f. envoi trop lourd : message dans l'écran, pas la page brute
limite = app.app.config["MAX_CONTENT_LENGTH"]
app.app.config["MAX_CONTENT_LENGTH"] = 2000
r = c.post(f"/etablissements/completer/excel?cle={ETAB}", content_type="multipart/form-data",
           data={"cle": ETAB, "fichiers": [(io.BytesIO(b"x" * 5000), "gros.xlsx")]})
app.app.config["MAX_CONTENT_LENGTH"] = limite
assert r.status_code == 302 and "trop lourds" in page(c.get(r.location)), r.status_code
print("ok  envoi au-delà de la limite : message « trop lourds », pas d'erreur brute")

# 10g. aperçu : la RH retouche une proposition avant d'enregistrer ; 422 garde les propositions
s = page(envoi(("lou.xlsx", xlsx([["Nom", "Prénom", "Sexe", "Date de naissance"],
                                  ["GARCIA", "Lou", "Feminin", "01/01/2001"]]))))
f = formulaire(s)
assert f[f"{GARCIA}:date_naissance"] == "01/01/2001" and f[f"{GARCIA}:civilite"] == "Madame"
f[f"{GARCIA}:date_naissance"] = "02/01/2001"                       # retouchée
f[f"{GARCIA}:date_debut"] = "pas une date"
r = c.post("/etablissements/completer", data=f)
s = page(r)
assert r.status_code == 422 and 'value="02/01/2001"' in s and "Madame</option>" in s, "propositions perdues au 422"
assert not store.lire(GARCIA)["champs"].get("civilite")
f[f"{GARCIA}:date_debut"] = ""
assert c.post("/etablissements/completer", data=f).status_code == 302
d = store.lire(GARCIA)["champs"]
assert d["date_naissance"] == "2001-01-02" and d["civilite"] == "Madame", "retouche ignorée"
print("ok  aperçu : retouche enregistrée, une erreur garde toutes les valeurs proposées")

# 10h. rattachements : même dossier deux fois, rattachement + saisie sur la même ligne
r = post(**{"x0": DARC, "x0:_nom": "A", "x0:telephone": "0611", "x1": DARC, "x1:_nom": "B", "x1:telephone": "0622"})
assert r.status_code == 422 and "déjà choisi" in page(r) and not store.lire(DARC)["champs"].get("telephone")
r = post(**{f"{DARC}:telephone": "0633", "x0": DARC, "x0:_nom": "A", "x0:telephone": "0611",
            "x0:adresse": "2 rue X"})
d = store.lire(DARC)["champs"]
assert r.status_code == 302 and d["telephone"] == "0633" and d["adresse"] == "2 rue X", d
print("ok  rattachement : même dossier deux fois refusé, la saisie du tableau prime sur la ligne rattachée")

# 10i. effacer un champ requis, date à un chiffre, session expirée
assert post(**{f"{DARC}:telephone": ""}).status_code == 302
s = page(c.get(f"/etablissements/completer?cle={ETAB}"))
assert store.lire(DARC)["champs"]["telephone"] == "" and "à compléter" in s
assert post(**{f"{DARC}:date_naissance": "1/2/1990"}).status_code == 302
assert store.lire(DARC)["champs"]["date_naissance"] == "1990-02-01"
expire = app.app.test_client()
r = expire.post("/etablissements/completer", data={"cle": ETAB, f"{DARC}:telephone": "0699"})
assert r.status_code == 302 and "/login" in r.location and store.lire(DARC)["champs"]["telephone"] == ""
r = expire.post(f"/etablissements/completer/excel?cle={ETAB}", content_type="multipart/form-data",
                data={"cle": ETAB, "fichiers": [(io.BytesIO(classeur), "a.xlsx")]})
assert r.status_code == 302 and "/login" in r.location
print("ok  champ requis effaçable (repasse « à compléter »), date à un chiffre, session expirée sans écriture")

print("\nOK test_excel")
